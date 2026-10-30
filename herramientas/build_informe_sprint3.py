# -*- coding: utf-8 -*-
"""Genera el Informe FINAL del Sprint 3 (docx) a partir de los resultados reales de la
evaluacion unica sobre la prueba oficial (HU12), el registro de modelos (HU14) y las
capturas reales de la aplicacion web (HU13).

Fuentes de verdad (no se inventa ningun numero):
  - "03. Implementacion/Sprint 3 - Evaluacion y despliegue/salida_sprint3/evaluacion_prueba.json"
  - "03. Implementacion/Sprint 3 - Evaluacion y despliegue/registro_modelos.json"
  - tarjetas de modelo en ".../modelos/*_tarjeta.json"
  - figuras ya generadas en "herramientas/figs_sprint3/"

Uso:
    python build_informe_sprint3.py   # arma herramientas/informe_sprint3.docx
"""
import json
import re
from pathlib import Path as _P

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

_AQUI = _P(__file__).resolve().parent
_RAIZ = _AQUI.parent
S3 = _RAIZ / "03. Implementacion" / "Sprint 3 - Evaluacion y despliegue"
FIGS = _AQUI / "figs_sprint3"
OUT = str(_AQUI / "informe_sprint3.docx")

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2F, 0x55, 0x97)
TEAL = RGBColor(0x2B, 0x7A, 0x8B)
GRAY = RGBColor(0x59, 0x59, 0x59)

# ================================================================ datos reales
EV = json.loads((S3 / "salida_sprint3" / "evaluacion_prueba.json").read_text(encoding="utf-8"))
REG = json.loads((S3 / "registro_modelos.json").read_text(encoding="utf-8"))
TARJ_CLF = json.loads((S3 / "modelos" / "clasificador_final_tarjeta.json").read_text(encoding="utf-8"))
TARJ_UNET = json.loads((S3 / "modelos" / "unet_final_tarjeta.json").read_text(encoding="utf-8"))

CLF = EV["clasificacion"]
SEG = EV["segmentacion"]
METAS = EV["metas_acta"]
MODELOS = {m["nombre"]: m for m in REG["modelos"]}
M_CLF = MODELOS["clasificador_final"]
M_UNET = MODELOS["unet_final"]

_ETIQUETA_CLASE = {"glioma": "Glioma", "meningioma": "Meningioma", "pituitary": "Hipófisis", "no_tumor": "Sin tumor"}
_ETIQUETA_PLANO = {"ax": "Axial", "co": "Coronal", "sa": "Sagital"}
_ETIQUETA_MODO = {"L": "Gris (L)", "RGB": "RGB"}

# recall minimo por clase (clase mas debil) — para el texto del resumen
_recall_min_clase = min(CLF["recall_por_clase"].items(), key=lambda kv: kv[1])
RECALL_MIN = _recall_min_clase[1] * 100
RECALL_MIN_CLASE = _ETIQUETA_CLASE[_recall_min_clase[0]]

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)

st = d.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.12
for name, size, color, before, after in (("Heading 1", 17, NAVY, 20, 8), ("Heading 2", 13, BLUE, 14, 5), ("Heading 3", 11, NAVY, 10, 3)):
    h = d.styles[name]
    h.font.name = "Calibri"
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h.element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    h.element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = color
    h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after)
    h.paragraph_format.keep_with_next = True


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def cell_margins(tbl, top=50, bottom=50, left=90, right=90):
    tblPr = tbl._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def borders(tbl, color="BFBFBF"):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def H1(t): return d.add_heading(t, 1)
def H2(t): return d.add_heading(t, 2)
def H3(t): return d.add_heading(t, 3)


def add_rich(p, text, bold=False, italic=False, size=None, color=None):
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color


def P(text="", bold=False, italic=False, size=None, color=None, align=None, after=None, keep=False):
    p = d.add_paragraph()
    add_rich(p, text, bold, italic, size, color)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if keep: p.paragraph_format.keep_with_next = True
    return p


def B(text, level=0):
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.63 + 0.6 * level)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(3)
    add_rich(p, text)
    return p


_fig_n = [0]


def T(rows, widths, header=True, size=9, zebra=True):
    tbl = d.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    borders(tbl); cell_margins(tbl)
    for j, wcm in enumerate(widths):
        tbl.columns[j].width = Cm(wcm)
    for i, row in enumerate(rows):
        tr = tbl.rows[i]
        for j, txt in enumerate(row):
            c = tr.cells[j]
            c.width = Cm(widths[j])
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            is_head = header and i == 0
            add_rich(p, str(txt), bold=is_head, size=size, color=RGBColor(255, 255, 255) if is_head else None)
            if is_head:
                shade(c, "1F3864")
            elif zebra and i % 2 == 0:
                shade(c, "F2F5FA")
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def FIG(path, width_cm, caption):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    c = d.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _fig_n[0] += 1
    caption = "Figura %d. " % _fig_n[0] + re.sub(r"^Figura \d+\. ", "", caption)
    add_rich(c, caption, italic=True, size=9, color=GRAY)
    c.paragraph_format.space_after = Pt(10)


def NOTE(text):
    t = d.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t, 90, 90, 160, 160)
    c = t.rows[0].cells[0]
    c.width = Cm(16.6)
    shade(c, "EAF0FA")
    tcPr = c._tc.get_or_add_tcPr()
    bd = OxmlElement("w:tcBorders")
    e = OxmlElement("w:left"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "24"); e.set(qn("w:color"), "2F5597"); bd.append(e)
    tcPr.append(bd)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    add_rich(p, text, size=10)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break():
    d.add_page_break()


def add_field(run, instr):
    for kind, txt in (("begin", None), (None, instr), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); run._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; run._r.append(e)


fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Informe Final del Sprint 3 · Sistema Inteligente: tumores cerebrales en MRI · Página ")
r.font.size = Pt(8.5); r.font.color.rgb = GRAY
r2 = fp.add_run(); r2.font.size = Pt(8.5); r2.font.color.rgb = GRAY
add_field(r2, "PAGE")
sec.different_first_page_header_footer = True

# ================================================================ PORTADA
for _ in range(3):
    d.add_paragraph()
band = d.add_table(rows=1, cols=1)
band.alignment = WD_TABLE_ALIGNMENT.CENTER
band.autofit = False
cell_margins(band, 300, 300, 300, 300)
c = band.rows[0].cells[0]
c.width = Cm(16.6)
shade(c, "1F3864")
p = c.paragraphs[0]
add_rich(p, "INFORME FINAL — SPRINT 3", bold=True, size=30, color=RGBColor(255, 255, 255))
p.paragraph_format.space_after = Pt(6)
p = c.add_paragraph()
add_rich(p, "Evaluación y despliegue", size=14, color=RGBColor(0xC9, 0xD6, 0xEC))
p.paragraph_format.space_after = Pt(0)

d.add_paragraph()
P("Desarrollo de un Sistema Inteligente para clasificar y segmentar tumores cerebrales en resonancias "
  "magnéticas (¿para qué?), desplegado en una aplicación web de apoyo académico (¿para dónde?)",
  bold=True, size=13, color=NAVY, after=14)

T([["Proyecto", "Proyecto académico de Redes Neuronales – Opción B (Ingeniería de Sistemas, Inteligencia Artificial)"],
   ["Metodología", "Híbrida: Scrum + CRISP-ML(Q) — fase de evaluación y despliegue"],
   ["Sprint", "Sprint 3 (24–30 oct 2026), presentación final 30 oct 2026"],
   ["Docente revisor", "Ing. Hugo Espetia Huamanga"],
   ["Dataset", "BRISC 2025 (Kaggle, CC BY 4.0) — prueba oficial de 1.000 imágenes (860 con máscara)"],
   ["Repositorio", "github.com/Calamar49/CNN-BRISC"]], [4.5, 12.1])

d.add_paragraph()
T([["Integrante", "Rol (propuesta)"],
   ["Coavoy Cruz Joseph Gabriel", "Gerente de Proyecto"],
   ["Cuchuyrrumi Mamani Manuel Rodrigo", "Scrum Master"],
   ["Huallpatuiro Rafaile Brayan", "Equipo de desarrollo"],
   ["Mamani Acuña Frank Joseph", "Equipo de desarrollo"]], [10, 6.6])

P("Este documento es una demostración académica; no es una herramienta de diagnóstico médico. Todos los "
  "números provienen de mediciones reales del equipo: la evaluación única de la prueba oficial está en "
  "`evaluacion_prueba.json` y el registro de modelos en `registro_modelos.json`. Las cifras del artículo de "
  "BRISC se citan como tales (otra partición), no como resultados propios.", italic=True, size=9.5, color=GRAY)
page_break()

# ================================================================ 1. RESUMEN EJECUTIVO
H1("1. Resumen ejecutivo")
P("El Sprint 2 dejó dos modelos elegidos y validados, pero medidos **solo sobre validación**; la prueba "
  "oficial (1.000 imágenes) nunca se había tocado. Este Sprint 3 cierra la parte técnica del proyecto con la "
  "fase de evaluación y despliegue de CRISP-ML(Q): se reentrenó la configuración elegida sobre "
  "entrenamiento+validación, se exportó a ONNX, se evaluó **una sola vez** sobre la prueba oficial (HU12), se "
  "construyó una aplicación web de demostración (HU13) y se dejó un registro de modelos versionado (HU14).")
B(f"**Clasificación (HU12):** accuracy global **{CLF['accuracy']*100:.2f} %** en la prueba oficial de "
  f"{EV['n_test_clasificacion']:,} imágenes; recall por clase entre {RECALL_MIN:.2f} % ({RECALL_MIN_CLASE}) y "
  f"100 %; F1 macro {CLF['f1_macro']*100:.2f} %.".replace(",", "."))
B(f"**Segmentación (HU12):** IoU global **{SEG['iou']*100:.2f} %** y Dice **{SEG['dice']*100:.2f} %** sobre "
  f"las {EV['n_test_segmentacion']} imágenes de prueba con máscara.")
B(f"**Las tres metas del acta se superan en la prueba oficial** (no solo en validación): accuracy "
  f"{CLF['accuracy']*100:.2f} % ≥ {METAS['accuracy']*100:.0f} %; recall mínimo {RECALL_MIN:.2f} % ≥ "
  f"{METAS['recall_por_clase']*100:.0f} % por clase; IoU {SEG['iou']*100:.2f} % ≥ {METAS['iou']*100:.0f} %.")
B(f"**El desempeño en prueba ≈ validación:** no hay colapso por sobreajuste. La U-Net incluso da "
  f"{SEG['iou']*100:.2f} % en prueba frente a {float(M_UNET['validacion_sprint2'].split('%')[0]):.2f} % en "
  "validación (Sprint 2); el clasificador queda en línea con su recall medio de validación "
  f"({float(M_CLF['validacion_sprint2'].split('%')[0]):.2f} %).")
B("**HU13 — aplicación web** de demostración funcionando (backend FastAPI + ONNX, visor de radiología), con "
  "capturas reales en la sección 5. **HU14 — MLOps básico:** registro de modelos versionado con tarjetas de "
  "modelo y procedencia (sección 6).")
NOTE("**Advertencia honesta (léase junto con cualquier cifra de prueba):** 113–123 de las 1.000 imágenes de "
     "prueba tienen un vecino casi idéntico en el entrenamiento oficial (hallazgo dHash/correlación del "
     "Sprint 1). Como el modelo final se reentrenó con train+val, la cifra de prueba **puede estar algo "
     "inflada**; se detalla en la sección 4.7 y en las limitaciones (sección 7).")
NOTE("**Qué falta (gestión del equipo, no de este informe):** cerrar la gestión Scrum del Sprint 3 (Backlog "
     "del Sprint, Informe de Estado, Minuta, Registro de Riesgos actualizado, Retrospectiva) y la "
     "presentación final (HU16) — quedan a cargo del equipo.")

# ================================================================ 2. METODOLOGIA Y ENTORNO
H1("2. Metodología y entorno")
H2("2.1 CRISP-ML(Q) — fase de evaluación y despliegue")
P("El Sprint 3 corresponde a las dos últimas fases de CRISP-ML(Q): **evaluación** (medir el modelo final "
  "contra datos que nunca participaron en ninguna decisión) y **despliegue** (dejar el modelo servido y "
  "versionado). La regla del proyecto es estricta: la prueba oficial no se usó para elegir arquitectura, "
  "hiperparámetros ni época; se evaluó una sola vez, con los modelos ONNX finales — los mismos que sirve la "
  "aplicación web.")
B("**Evaluación (HU12):** una única pasada sobre la prueba oficial, reportando las metas del acta y el "
  "desglose por clase, plano y tipo de archivo (riesgo R09).")
B("**Despliegue (HU13):** backend local que sirve los modelos ONNX con `onnxruntime` e interfaz web.")
B("**Calidad y reproducibilidad (HU14):** registro de modelos versionado, tarjetas de modelo con procedencia "
  "y paridad, y commit de Git; los pesos y los ONNX no se versionan, se regeneran con los scripts.")

H2("2.2 ¿Por qué se reentrena con train+val antes de HU12?")
P("En el Sprint 2 los modelos se entrenaron solo con entrenamiento (3.999 imágenes) y se usó validación "
  "(1.001) para elegir arquitectura, hiperparámetros y número de épocas. Para el modelo final, siguiendo la "
  "buena práctica de aprendizaje automático, se **reentrenó la configuración ya elegida** sobre "
  "entrenamiento+validación juntos (más datos → mejor modelo), **fijando el número de épocas** (sin parada "
  "temprana, porque ya no hay un conjunto de validación separado del que leer la mejor época). Las épocas se "
  "fijaron como el promedio redondeado de las épocas óptimas de las 3 semillas del Sprint 2: clasificador "
  "6/14/10 → 10; U-Net 28/25/24 → 26. Así, la prueba oficial queda como un conjunto verdaderamente no visto "
  "para HU12.")

H2("2.3 Entorno de ejecución")
T([["Variable", "Valor"],
   ["Sistema operativo", "Windows 11"],
   ["Python", "3.12.8"],
   ["GPU", "NVIDIA GeForce RTX 5060 Laptop (8 GB)"],
   ["PyTorch", "2.11.0+cu128 (con CUDA)"],
   ["Exportación / servicio", "ONNX (opset 17) + onnxruntime"],
   ["Precisión (entrenamiento)", "mixta (torch.autocast)"]], [6, 10.6])

H2("2.4 Datos de esta fase")
P(f"El modelo final se reentrenó con entrenamiento+validación: 5.000 imágenes de clasificación "
  "(3.999 + 1.001) y 3.933 con máscara para segmentación (3.146 + 787; «sin tumor» no tiene máscara y queda "
  f"fuera de esa tarea). La evaluación de HU12 usa la **prueba oficial**: {EV['n_test_clasificacion']:,} "
  f"imágenes de clasificación y {EV['n_test_segmentacion']} con máscara para segmentación. Normalización de "
  "ImageNet, imagen en gris repetida en 3 canales, 256×256 — idéntico pipeline al del Sprint 2.".replace(",", "."))
page_break()

# ================================================================ 3. MODELO FINAL Y ONNX
H1("3. Modelo final y exportación a ONNX")
P("El modelo final de cada tarea es el reentrenamiento (train+val, épocas fijadas) de la configuración "
  "elegida en el Sprint 2. Luego se exportó a ONNX (opset 17) y se verificó la **paridad PyTorch↔ONNX** sobre "
  "imágenes reales, para garantizar que lo que sirve la aplicación web es numéricamente equivalente al "
  "modelo entrenado en PyTorch.")

H2("3.1 Configuración del reentrenamiento final")
T([["Modelo", "Arquitectura", "Datos de reentrenamiento", "Épocas (fijadas)", "lr", "batch"],
   ["Clasificador", "EfficientNetB0 (TL ImageNet)", "train+val, 5.000 imágenes",
    f"{M_CLF['entrenamiento']['epocas']} (6/14/10→10)", "3e-4", str(M_CLF['entrenamiento']['batch'])],
   ["U-Net", "ResNet34 (TL ImageNet)", "train+val, 3.933 con máscara",
    f"{M_UNET['entrenamiento']['epocas']} (28/25/24→26)", "1e-4", str(M_UNET['entrenamiento']['batch'])]],
  [2.6, 3.8, 4.3, 2.9, 1.5, 1.5])

H2("3.2 Exportación a ONNX y paridad")
T([["Modelo", "Entrada", "Salida", "Tamaño ONNX", "Paridad PyTorch↔ONNX"],
   ["Clasificador", "float32 [N,3,256,256]", "logits [N,4]", f"{TARJ_CLF['tamano_mb']:.1f} MB",
    "dif. máx. 4.1e-04; clases idénticas"],
   ["U-Net", "float32 [N,3,256,256]", "logits [N,1,256,256]", f"{TARJ_UNET['tamano_mb']:.1f} MB",
    "dif. máx. 3.45e-04; máscara coincide 100 % de píxeles"]],
  [2.4, 3.6, 3.4, 2.4, 4.8])
P("La diferencia máxima entre la salida de PyTorch y la de ONNX es del orden de 1e-04 en ambos modelos: las "
  "clases predichas por el clasificador son idénticas y la decisión de máscara de la U-Net (umbral 0,5) "
  "coincide en el 100 % de los píxeles. La exportación a ONNX no cambia, en la práctica, el resultado; "
  "habilita servir el modelo con `onnxruntime` sin depender de PyTorch en producción.", size=9.5, color=GRAY)
page_break()

# ================================================================ 4. HU12 EVALUACION PRUEBA
H1("4. HU12 — Evaluación única sobre la prueba oficial")
P("Esta es la única evaluación del proyecto sobre la prueba oficial. Se corrió una sola vez, con los modelos "
  "ONNX finales, y sus resultados **no retroalimentaron ninguna decisión** de diseño. Todas las cifras de "
  "esta sección provienen de `salida_sprint3/evaluacion_prueba.json`.")

H2("4.1 Resultado global y metas del acta")
T([["Tarea", "Métrica", "Prueba oficial (HU12)", "Meta del acta", "Validación (Sprint 2)"],
   ["Clasificación", "Accuracy global", f"{CLF['accuracy']*100:.2f} %", f"≥ {METAS['accuracy']*100:.0f} %",
    M_CLF['validacion_sprint2'].replace(" (3 semillas, Sprint 2)", "") + " (recall medio)"],
   ["Clasificación", "Recall mínimo por clase", f"{RECALL_MIN:.2f} % ({RECALL_MIN_CLASE})",
    f"≥ {METAS['recall_por_clase']*100:.0f} %", "—"],
   ["Clasificación", "F1 macro", f"{CLF['f1_macro']*100:.2f} %", "—", "—"],
   ["Segmentación", "IoU global", f"{SEG['iou']*100:.2f} %", f"≥ {METAS['iou']*100:.0f} %",
    M_UNET['validacion_sprint2'].replace(" (3 semillas, Sprint 2)", "") + " (IoU)"],
   ["Segmentación", "Dice global", f"{SEG['dice']*100:.2f} %", "—", "—"]],
  [2.6, 3.2, 3.5, 2.3, 5.0])
P("**Las tres metas del acta se superan en la prueba oficial**, no solo en validación. El margen es amplio: "
  f"accuracy {CLF['accuracy']*100:.2f} % frente al {METAS['accuracy']*100:.0f} % pedido, recall mínimo "
  f"{RECALL_MIN:.2f} % frente al {METAS['recall_por_clase']*100:.0f} %, e IoU {SEG['iou']*100:.2f} % frente "
  f"al {METAS['iou']*100:.0f} %.")

H2("4.2 Validación (Sprint 2) frente a prueba oficial")
P("La comparación clave de la evaluación: si el desempeño se desplomara en la prueba respecto de validación, "
  "sería señal de sobreajuste. No ocurre. El clasificador queda en línea con su recall medio de validación y "
  "la U-Net incluso mejora levemente en prueba, lo que es consistente con haber reentrenado el modelo final "
  "con más datos (train+val).")
FIG(FIGS / "val_vs_prueba.png", 15, "Validación (promedio de 3 semillas, Sprint 2) frente a prueba oficial "
    "(HU12), para el clasificador y la U-Net, con las líneas de meta del acta. El desempeño en prueba "
    "acompaña al de validación en ambas tareas; no hay colapso por sobreajuste.")

H2("4.3 Matriz de confusión en la prueba")
P(f"Matriz de confusión real sobre las {EV['n_test_clasificacion']:,} imágenes de la prueba oficial. Los "
  "pocos errores se concentran en la frontera glioma↔meningioma; «hipófisis» y «sin tumor» no tienen ningún "
  "error.".replace(",", "."))
FIG(FIGS / "prueba_matriz_confusion.png", 10, "Matriz de confusión real sobre las 1.000 imágenes de la "
    "prueba oficial (modelos ONNX finales). Las clases «hipófisis» y «sin tumor» se predicen sin error.")

H2("4.4 Recall por clase en la prueba")
P("El recall por clase es la métrica más exigente del acta (≥ 90 % en cada clase). Las cuatro clases lo "
  "superan con holgura; glioma es la más baja, coherente con su bajo contraste y con que también es la clase "
  "más difícil de segmentar (sección 4.5).")
FIG(FIGS / "prueba_recall_por_clase.png", 14, "Recall por clase en la prueba oficial, con la línea de meta "
    "del 90 %. Las cuatro clases la superan; glioma es la más baja (98,03 %).")
rows = [["Clase", "Recall en prueba", "Meta del acta"]]
for k in ("glioma", "meningioma", "pituitary", "no_tumor"):
    rows.append([_ETIQUETA_CLASE[k], f"{CLF['recall_por_clase'][k]*100:.2f} %", f"≥ {METAS['recall_por_clase']*100:.0f} %"])
T(rows, [5.5, 5.5, 5.6])

H2("4.5 IoU y Dice por clase de tumor")
P("La segmentación se reporta por clase (solo las 3 clases con tumor tienen máscara). El promedio global de "
  f"{SEG['iou']*100:.2f} % de IoU oculta que glioma es, con claridad, la clase más difícil, igual que en el "
  "Sprint 2: su contorno es el menos definido por su bajo contraste con el tejido sano.")
FIG(FIGS / "prueba_iou_por_clase.png", 14, "IoU y Dice por clase de tumor en la prueba oficial. Glioma es la "
    "clase más difícil de segmentar; meningioma la más fácil.")
rows = [["Clase", "IoU en prueba", "Dice en prueba", "n (imágenes)"]]
for k in ("glioma", "meningioma", "pituitary"):
    rows.append([_ETIQUETA_CLASE[k], f"{SEG['iou_por_clase'][k]['valor']*100:.2f} %",
                 f"{SEG['dice_por_clase'][k]['valor']*100:.2f} %", str(SEG['iou_por_clase'][k]['n'])])
T(rows, [4.6, 4.0, 4.0, 4.0])

H2("4.6 Desglose por plano de corte y por tipo de archivo (riesgo R09)")
P("El Sprint 1 advirtió un posible atajo por procedencia: el tipo de archivo de origen (gris/RGB) cambia con "
  "la clase. Se repitió en la prueba oficial el desglose por plano de corte y por tipo de archivo para ambas "
  "tareas, el análisis que quedó pendiente del Sprint 2 (riesgo R09).")
rows = [["Plano", "Accuracy clasif.", "n (clasif.)", "IoU U-Net", "n (seg.)"]]
for k in ("ax", "co", "sa"):
    rows.append([_ETIQUETA_PLANO[k], f"{CLF['accuracy_por_plano'][k]['valor']*100:.2f} %",
                 str(CLF['accuracy_por_plano'][k]['n']), f"{SEG['iou_por_plano'][k]['valor']*100:.2f} %",
                 str(SEG['iou_por_plano'][k]['n'])])
T(rows, [3.2, 3.6, 3.0, 3.4, 3.0])
FIG(FIGS / "prueba_por_modo.png", 14, "Accuracy (clasificador) e IoU (U-Net) por tipo de archivo de origen "
    "(gris/RGB) en la prueba oficial — eje del riesgo R09.")
rows = [["Tipo de archivo", "Accuracy clasif.", "n (clasif.)", "IoU U-Net", "n (seg.)"]]
for k in ("L", "RGB"):
    rows.append([_ETIQUETA_MODO[k], f"{CLF['accuracy_por_modo'][k]['valor']*100:.2f} %",
                 str(CLF['accuracy_por_modo'][k]['n']), f"{SEG['iou_por_modo'][k]['valor']*100:.2f} %",
                 str(SEG['iou_por_modo'][k]['n'])])
T(rows, [3.6, 3.6, 3.0, 3.4, 3.0])
_brecha_modo = (SEG['iou_por_modo']['RGB']['valor'] - SEG['iou_por_modo']['L']['valor']) * 100
B(f"**Clasificador:** sin diferencias relevantes por plano ({CLF['accuracy_por_plano']['sa']['valor']*100:.2f} "
  f"%–{CLF['accuracy_por_plano']['ax']['valor']*100:.2f} %). Por tipo de archivo, gris "
  f"{CLF['accuracy_por_modo']['L']['valor']*100:.2f} % vs. RGB {CLF['accuracy_por_modo']['RGB']['valor']*100:.2f} "
  "% — una brecha mínima.")
B(f"**U-Net:** la brecha gris ({SEG['iou_por_modo']['L']['valor']*100:.2f} %) vs. RGB "
  f"({SEG['iou_por_modo']['RGB']['valor']*100:.2f} %) es de {_brecha_modo:.1f} puntos de IoU, más chica que la "
  "del Sprint 2 en validación (~8 pp). Como el glioma —la clase más difícil— es mayoritariamente en escala "
  "de grises (88 %, Sprint 1), la brecha parece explicarse en buena parte por la dificultad propia del "
  "glioma y no por un atajo puro de procedencia. R09 queda documentado, no concluido (sección 7).")

H2("4.7 Advertencia sobre posible fuga en la prueba")
NOTE("**Lea esto junto con toda cifra de prueba.** El Sprint 1 encontró que **113–123 de las 1.000 imágenes "
     "de prueba** tienen un vecino casi idéntico en el entrenamiento oficial (hallazgo dHash/correlación). "
     "Como el modelo final se reentrenó con train+val, parte de esas imágenes de prueba se parecen mucho a "
     "imágenes vistas en el entrenamiento, así que **la cifra de prueba puede estar algo inflada**. Esto no "
     "invalida que las metas se superen, pero obliga a leer los resultados con esa reserva; se recoge en las "
     "limitaciones (sección 7).")
P(f"Para contexto (no es nuestro resultado): el artículo de BRISC reporta, con otra partición, EfficientNetB0 "
  f"99,2 % de accuracy y U-Net 75,7 % de mIoU. Nuestras cifras de prueba ({CLF['accuracy']*100:.2f} % y "
  f"{SEG['iou']*100:.2f} %) están en ese orden o por encima, pero no son comparables en igualdad de "
  "condiciones (otra partición y otro protocolo).", size=9.5, color=GRAY)
page_break()

# ================================================================ 5. HU13 APP WEB
H1("5. HU13 — Aplicación web de demostración")
P("Se construyó una aplicación web que sirve los modelos ONNX finales: se sube una resonancia y la app "
  "predice el tipo de tumor y marca su contorno. El **backend** es FastAPI con `onnxruntime`; la **interfaz** "
  "es web vanilla (HTML/CSS/JS) con un diseño de «visor de radiología»: imagen en gris sobre fondo oscuro, "
  "máscara de tumor como *overlay* coral y lecturas en tipografía monoespaciada. Trae ejemplos reales "
  "tomados de la prueba oficial.")

H2("5.1 Decisión de arquitectura (riesgo R05)")
P("Se eligió **backend local** (el modelo corre en el servidor, no en el navegador) por el riesgo R05: la "
  f"U-Net pesa {TARJ_UNET['tamano_mb']:.1f} MB en ONNX, demasiado para descargarla y ejecutarla cómodamente "
  "en el navegador del usuario. El clasificador ONNX pesa "
  f"{TARJ_CLF['tamano_mb']:.1f} MB. Con el backend local, la app carga rápido y el usuario solo envía la "
  "imagen y recibe la predicción.")

H2("5.2 Interfaz en funcionamiento (capturas reales)")
P("Capturas reales de la aplicación corriendo con los modelos ONNX finales.")
FIG(FIGS / "app_vacio.jpg", 14, "Interfaz de la aplicación antes de cargar una imagen: visor oscuro, zona de "
    "carga y galería de ejemplos reales de la prueba oficial.")
FIG(FIGS / "app_glioma.jpg", 14, "Caso real de glioma: el clasificador predice glioma con 100 % de confianza "
    "y la U-Net marca el contorno del tumor (área marcada 2,24 % de la imagen).")
FIG(FIGS / "app_meningioma.jpg", 14, "Caso real de meningioma: 99,1 % de confianza, área marcada 4,47 % de la "
    "imagen; el contorno coral se superpone a la resonancia en gris.")
FIG(FIGS / "app_sin_tumor.jpg", 14, "Caso real sin tumor: el clasificador predice «sin tumor» con 100 % de "
    "confianza. La segmentación se muestra solo como información; cuando el clasificador dice «sin tumor», la "
    "máscara no debe tomarse como un hallazgo.")
NOTE("La aplicación es una **demostración académica**, no una herramienta de diagnóstico. Cuando el "
     "clasificador predice «sin tumor», la máscara de segmentación es meramente informativa: la U-Net se "
     "entrenó solo con imágenes que tienen tumor y nunca aprende a afirmar la ausencia de tumor.")
page_break()

# ================================================================ 6. HU14 MLOPS
H1("6. HU14 — MLOps básico")
P("El despliegue va acompañado de un registro de modelos versionado, para que cualquier integrante pueda "
  "saber qué modelo está sirviendo la app, de dónde salió y cómo regenerarlo. Dos piezas: una **tarjeta de "
  "modelo** por artefacto (`*_tarjeta.json`) y un **registro** central (`registro_modelos.json`) que sí se "
  "sube a Git.")

H2("6.1 Registro de modelos")
T([["Modelo", "Formato", "Tamaño", "Validación (Sprint 2)", "Prueba oficial (HU12)", "Commit"],
   ["clasificador_final", TARJ_CLF['formato'], f"{TARJ_CLF['tamano_mb']:.1f} MB",
    M_CLF['validacion_sprint2'].replace(" (3 semillas, Sprint 2)", ""),
    f"accuracy {M_CLF['prueba_oficial_hu12']['accuracy']:.2f} %", REG['commit_git']],
   ["unet_final", TARJ_UNET['formato'], f"{TARJ_UNET['tamano_mb']:.1f} MB",
    M_UNET['validacion_sprint2'].replace(" (3 semillas, Sprint 2)", ""),
    f"IoU {M_UNET['prueba_oficial_hu12']['iou']:.2f} %", REG['commit_git']]],
  [3.2, 2.6, 1.9, 3.4, 3.1, 2.4])

H2("6.2 Tarjetas de modelo")
P("Cada modelo lleva una tarjeta con su formato, tamaño, forma de entrada/salida, la paridad PyTorch↔ONNX y "
  "la procedencia (datos, épocas, hiperparámetros, semilla, optimizador). Resumen de ambas tarjetas:")
T([["Campo", "clasificador_final", "unet_final"],
   ["Formato", TARJ_CLF['formato'], TARJ_UNET['formato']],
   ["Entrada", "float32 [N,3,256,256]", "float32 [N,3,256,256]"],
   ["Salida", "logits [N,4]", "logits [N,1,256,256] (sigmoide + umbral 0,5)"],
   ["Datos", "train+val", "train+val (solo con máscara)"],
   ["Épocas / lr / batch", "10 / 3e-4 / 64", "26 / 1e-4 / 8"],
   ["Paridad PyTorch↔ONNX", "dif. máx. 4.14e-04", "dif. máx. 3.45e-04"]],
  [4.0, 5.0, 7.6])

H2("6.3 Qué se versiona y qué no")
B("**Sí se sube a Git:** el registro `registro_modelos.json` (con procedencia, métrica de validación y de "
  "prueba, y el commit de Git), las tarjetas de modelo `*_tarjeta.json` y todos los scripts "
  "(`reentrenar_final.py`, `exportar_onnx.py`, `evaluar_prueba.py`, `registrar_modelos.py`).")
B("**No se sube a Git:** los pesos (`.pt`) ni los modelos (`.onnx`), por tamaño y por la regla del proyecto; "
  "se regeneran de forma reproducible con los scripts a partir del dataset preparado.")
page_break()

# ================================================================ 7. LIMITACIONES
H1("7. Limitaciones y amenazas a la validez")
P("Las metas se superan, pero varios resultados tienen un alcance más chico del que parece a primera lectura. "
  "Esta sección reúne, en un solo lugar, qué tan lejos se puede estirar cada cifra.")
B("**La prueba se evaluó una sola vez:** por diseño del proyecto, no hay repeticiones ni intervalos de "
  "confianza sobre la prueba oficial; es una única medición, no una distribución.")
B("**Posible fuga (113–123 imágenes):** entre 113 y 123 de las 1.000 imágenes de prueba tienen un vecino "
  "casi idéntico en el entrenamiento oficial; como el modelo final se reentrenó con train+val, la cifra de "
  "prueba puede estar algo inflada. Es la reserva más importante al leer este informe (sección 4.7).")
B(f"**Glioma es la clase más difícil en las dos tareas:** el recall más bajo de clasificación es glioma "
  f"({CLF['recall_por_clase']['glioma']*100:.2f} %) y el IoU más bajo de segmentación también es glioma "
  f"({SEG['iou_por_clase']['glioma']['valor']*100:.2f} %), coherente con su bajo contraste con el tejido "
  "sano (hallazgo del Sprint 1). Es la clase a vigilar en cualquier uso futuro.")
B("**R09 queda documentado, no concluido:** la brecha gris/RGB en segmentación se midió en la prueba, pero "
  "sigue mezclada con el hecho de que el glioma es mayoritariamente en escala de grises; separar cuánto es "
  "atajo de procedencia y cuánto es dificultad propia del glioma necesitaría más datos de los disponibles.")
B("**Alcance del dataset:** BRISC 2025 es un conjunto público ya curado, de un solo tipo de secuencia (T1 "
  "con contraste); no se probó con otro dataset, escáner ni secuencia, así que estas cifras no dicen cómo "
  "generalizaría el modelo a resonancias de otra procedencia.")
B("**La demostración no es diagnóstico:** la app es una prueba de concepto académica. La U-Net nunca aprende "
  "a afirmar la ausencia de tumor (se entrenó solo con imágenes con máscara), por lo que cuando el "
  "clasificador dice «sin tumor» la máscara es solo informativa; un sistema real debería combinar ambas "
  "salidas con cuidado y pasar por validación clínica.")

# ================================================================ 8. CONCLUSIONES
H1("8. Conclusiones y cierre del proyecto")
H2("8.1 Conclusiones")
B(f"La evaluación única sobre la prueba oficial (HU12) **confirma las tres metas del acta**: accuracy "
  f"{CLF['accuracy']*100:.2f} %, recall mínimo {RECALL_MIN:.2f} % por clase e IoU {SEG['iou']*100:.2f} %.")
B("El desempeño en prueba acompaña al de validación: no hay colapso por sobreajuste, y el reentrenamiento con "
  "train+val fijando las épocas funcionó como se esperaba.")
B("El modelo final se exportó a ONNX con paridad verificada (dif. máx. ~1e-04, decisión idéntica), lo que "
  "permitió servirlo desde una aplicación web de demostración (HU13) con backend local por el riesgo R05.")
B("Queda un registro de modelos versionado con tarjetas y procedencia (HU14): el proyecto es reproducible y "
  "trazable sin subir pesos ni modelos a Git.")
B("La reserva honesta del proyecto se mantiene explícita: la posible fuga de 113–123 imágenes de prueba y la "
  "dificultad del glioma acompañan cualquier lectura de estos resultados.")

H2("8.2 Riesgos del proyecto (cierre)")
T([["Riesgo", "Descripción", "Estado / mitigación"],
   ["R01", "Fuga entre pacientes", "Mitigado: partición propia por `grupo_hash`; se reporta la advertencia de las 113–123 imágenes"],
   ["R04", "Arquitecturas inestables", "Mitigado: 2 arquitecturas y 3 semillas (Sprint 2)"],
   ["R05", "Modelo pesado para el navegador", "Mitigado: backend local + ONNX (U-Net 97,7 MB)"],
   ["R06", "Memoria de GPU 8 GB", "Mitigado: precisión mixta y lote pequeño"],
   ["R09", "Atajo por procedencia gris/RGB", "Investigado en la prueba por tipo de archivo; documentado, no concluido"]],
  [1.6, 5.2, 9.8])

H2("8.3 Qué falta (gestión del equipo)")
P("Por fuera de lo técnico, el cierre del Sprint 3 todavía necesita, a cargo del equipo y no de este informe: "
  "el Backlog del Sprint 3, el Informe de Estado y la Minuta actualizados, el Registro de Riesgos con el "
  "estado final de R01–R09, la Retrospectiva del Sprint 3 (con el equipo) y la presentación final (HU16).")
page_break()

# ================================================================ 9. GLOSARIO
H1("9. Glosario mínimo")
B("**ONNX (*Open Neural Network Exchange*):** formato estándar para guardar un modelo entrenado de forma "
  "independiente del framework; permite servirlo con `onnxruntime` sin depender de PyTorch en producción.")
B("**IoU (*Intersection over Union*):** área donde la máscara predicha y la real coinciden, dividida entre el "
  "área que ocupan entre las dos juntas; 100 % es coincidencia perfecta.")
B("**Dice:** mide lo mismo que el IoU (el solapamiento entre máscara predicha y real), con una fórmula que da "
  "más peso al área que sí coincide; el Dice de una misma predicción siempre es igual o mayor que su IoU.")
B("**Recall:** de los casos reales de una clase, qué porcentaje detecta el modelo (por ejemplo, de los "
  "gliomas reales de la prueba, cuántos el modelo etiquetó como glioma).")
B("**Transfer learning:** partir de una red ya entrenada con otro conjunto (aquí, ImageNet) en vez de empezar "
  "con pesos al azar, para aprovechar lo que esa red ya «sabe» sobre bordes, texturas y formas.")
B("**Fuga (*data leakage*):** que imágenes casi idénticas caigan en el entrenamiento y en el conjunto con el "
  "que se mide el resultado, lo que puede inflar la métrica sin que el modelo generalice mejor de verdad.")
B("**MLOps:** prácticas para versionar, registrar y servir modelos de forma reproducible y trazable (aquí: "
  "registro de modelos, tarjetas de modelo y procedencia, exportación a ONNX).")
B("**Paridad PyTorch↔ONNX:** comprobar que el modelo exportado a ONNX produce, sobre imágenes reales, "
  "prácticamente la misma salida que el modelo original en PyTorch (misma clase, misma máscara).")

d.save(OUT)
print("Informe guardado en:", OUT)
