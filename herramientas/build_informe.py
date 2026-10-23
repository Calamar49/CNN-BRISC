# -*- coding: utf-8 -*-
"""Genera el Informe del Sprint 1 (docx) a partir de los resultados reales guardados en salida_sprint1/."""
from pathlib import Path as _P
_AQUI = _P(__file__).resolve().parent
_RAIZ = _AQUI.parent
import copy
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PROY = str(_RAIZ).replace("\\", "/") + "/"
SAL = PROY + "03. Implementacion/Sprint 1 - Notebooks/salida_sprint1/"
FIGS = str(_AQUI / "figs").replace("\\", "/") + "/"
_fig_n = [0]
OUT = str(_AQUI / "informe_sprint1.docx")

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2F, 0x55, 0x97)
GRAY = RGBColor(0x59, 0x59, 0x59)

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)

st = d.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.12
for name, size, color, before, after in (("Heading 1", 17, NAVY, 20, 8), ("Heading 2", 13, BLUE, 14, 5), ("Heading 3", 11, NAVY, 10, 3)):
    h = d.styles[name]
    h.font.name = "Calibri"
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h.element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    h.element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = color
    h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after)
    h.paragraph_format.keep_with_next = True


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def cell_margins(tbl, top=50, bottom=50, left=90, right=90):
    tblPr = tbl._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def borders(tbl, color="BFBFBF"):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def H1(t): return d.add_heading(t, 1)
def H2(t): return d.add_heading(t, 2)
def H3(t): return d.add_heading(t, 3)


def P(text="", bold=False, italic=False, size=None, color=None, align=None, after=None, keep=False):
    p = d.add_paragraph()
    add_rich(p, text, bold, italic, size, color)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if keep: p.paragraph_format.keep_with_next = True
    return p


def add_rich(p, text, bold=False, italic=False, size=None, color=None):
    """**negrita** dentro del texto"""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part: continue
        segs = part.split("`")
        for k, seg in enumerate(segs):
            if not seg: continue
            r = p.add_run(seg)
            r.bold = bold or (i % 2 == 1)
            r.italic = italic
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
            if k % 2 == 1:
                r.font.name = "Consolas"
                r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
                r.font.size = Pt((size or 10.5) - 1)


def B(text, level=0):
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.63 + 0.6 * level)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(3)
    add_rich(p, text)
    return p


def N(text):
    p = d.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    add_rich(p, text)
    return p


def T(rows, widths, header=True, size=9, align_right=(), zebra=True, bold_last=False):
    tbl = d.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    borders(tbl); cell_margins(tbl)
    for j, wcm in enumerate(widths):
        tbl.columns[j].width = Cm(wcm)
    for i, row in enumerate(rows):
        tr = tbl.rows[i]
        trPr = tr._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        if i == 0 and header:
            th = OxmlElement("w:tblHeader"); trPr.append(th)
        for j, txt in enumerate(row):
            c = tr.cells[j]
            c.width = Cm(widths[j])
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if i < len(rows) - 1 or True:
                p.paragraph_format.keep_with_next = i < len(rows) - 1 and len(rows) <= 14
            if j in align_right:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            is_head = header and i == 0
            add_rich(p, str(txt), bold=is_head or (bold_last and i == len(rows) - 1), size=size,
                     color=RGBColor(255, 255, 255) if is_head else None)
            if is_head:
                shade(c, "1F3864")
            elif zebra and i % 2 == 0:
                shade(c, "F2F5FA")
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def FIG(path, width_cm, caption):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = d.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _fig_n[0] += 1
    caption = "Figura %d. " % _fig_n[0] + re.sub(r"^Figura \d+\. ", "", caption)
    add_rich(c, caption, italic=True, size=9, color=GRAY)
    c.paragraph_format.space_after = Pt(10)


def NOTE(text):
    t = d.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t, 90, 90, 160, 160)
    c = t.rows[0].cells[0]
    c.width = Cm(16.6)
    shade(c, "EAF0FA")
    tcPr = c._tc.get_or_add_tcPr()
    bd = OxmlElement("w:tcBorders")
    e = OxmlElement("w:left"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "24"); e.set(qn("w:color"), "2F5597"); bd.append(e)
    tcPr.append(bd)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    add_rich(p, text, size=10)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break():
    d.add_page_break()


# ---------------------------------------------------------------- pie de pagina
def add_field(run, instr):
    for kind, txt in (("begin", None), (None, instr), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); run._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; run._r.append(e)


fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Informe del Sprint 1 · Sistema Inteligente: tumores cerebrales en MRI · Página ")
r.font.size = Pt(8.5); r.font.color.rgb = GRAY
r2 = fp.add_run(); r2.font.size = Pt(8.5); r2.font.color.rgb = GRAY
add_field(r2, "PAGE")
sec.different_first_page_header_footer = True

# ================================================================ PORTADA
for _ in range(3):
    d.add_paragraph()
band = d.add_table(rows=1, cols=1)
band.alignment = WD_TABLE_ALIGNMENT.CENTER
band.autofit = False
cell_margins(band, 300, 300, 300, 300)
c = band.rows[0].cells[0]
c.width = Cm(16.6)
shade(c, "1F3864")
p = c.paragraphs[0]
add_rich(p, "INFORME DEL SPRINT 1", bold=True, size=30, color=RGBColor(255, 255, 255))
p.paragraph_format.space_after = Pt(6)
p = c.add_paragraph()
add_rich(p, "Comprensión del negocio y de los datos · Preparación de los datos", size=14, color=RGBColor(0xF6, 0xA8, 0x1C))
p.paragraph_format.space_after = Pt(0)

d.add_paragraph()
P("Desarrollo de un Sistema Inteligente para clasificar y segmentar tumores cerebrales en resonancias magnéticas (¿para qué?), "
  "desplegado en una aplicación web de apoyo académico (¿para dónde?)", bold=True, size=14, color=NAVY, after=14)

T([["Proyecto", "Proyecto académico de Redes Neuronales – Opción B (Ingeniería de Sistemas, Inteligencia Artificial)"],
   ["Metodología", "Híbrida: Scrum + CRISP-ML(Q)"],
   ["Dataset", "BRISC 2025 (6.000 resonancias T1 con contraste; clasificación y segmentación)"],
   ["Periodo del Sprint 1", "Trabajo del 1 al 8 de octubre de 2026 · Presentación al docente: lunes 12 de octubre de 2026"],
   ["Fecha del informe", "14 de octubre de 2026 (versión ampliada, con el capítulo de comprensión del negocio)"],
   ["Equipo", "Coavoy Cruz Joseph Gabriel (Gerente de Proyecto) · Cuchuyrrumi Mamani Manuel Rodrigo (Scrum Master) · "
              "Huallpatuiro Rafaile Brayan · Mamani Acuña Frank Joseph (Equipo de desarrollo)"],
   ["Docente revisor", "Ing. Hugo Espetia Huamanga"],
   ["Repositorio", "https://github.com/Calamar49/CNN-BRISC"],
   ["Tablero Scrum", "https://trello.com/b/swXjkoyS/product-backlog-proyecto-redes-neuronales-brisc"]],
  [3.6, 13.0], header=False, size=9.5, zebra=False)
P("Es una demostración académica: no es una herramienta de diagnóstico médico.", italic=True, size=9.5, color=GRAY)
page_break()

# ================================================================ CONTENIDO
H1("Contenido")
for t in ["1. Resumen ejecutivo", "2. Comprensión del negocio", "    2.1 El problema y su contexto", "    2.2 El sector privado de salud (datos públicos)", "    2.3 A quién está dirigido", "    2.4 Escenario de referencia: clínica privada con resonancia", "    2.5 a 2.11 Objetivos, situación, factibilidad, ética, metodología y equipo", "3. Plan del Sprint 1", "4. Trabajo realizado",
          "    4.1 HU01 · Definir el problema y el alcance", "    4.2 HU02 · Descargar y explorar BRISC (comprensión de los datos)",
          "    4.3 HU03 · Preparar los datos y definir la partición", "    4.4 HU04 · Presentación del Sprint 1",
          "5. Hallazgos y decisiones técnicas", "6. Calidad y verificación", "7. Cambios respecto al plan inicial",
          "8. Riesgos y limitaciones", "9. Entorno de trabajo y reproducibilidad", "10. Alcance: lo que no se hizo en este sprint",
          "11. Plan del Sprint 2 y recomendaciones", "Anexo A. Inventario de entregables", "Anexo B. Referencias"]:
    P(t, after=1)
page_break()

# ================================================================ 1. RESUMEN
H1("1. Resumen ejecutivo")
P("El Sprint 1 cubrió las dos primeras fases de CRISP-ML(Q): **comprensión del negocio y de los datos** y **preparación de los datos**. "
  "Se completaron las 4 historias de usuario planificadas (HU01 a HU04, 14 puntos). El Acta de Constitución está firmada y el equipo trabaja con "
  "un tablero Scrum de 17 historias para todo el proyecto.")
P("Resultados principales:", bold=True, keep=True)
B("**Comprensión del negocio:** se describió el problema con datos públicos (Perú: 1.616 casos nuevos de cáncer de cerebro y sistema nervioso central en 2024 según GLOBOCAN; 495 radiólogos, el 54,8 % de radiólogos y tecnólogos en Lima), a quién iría dirigido el sistema, un escenario hipotético de centro privado de diagnóstico por imágenes con supuestos marcados como tales, objetivos, factibilidad y consideraciones éticas. El proyecto no tiene cliente real.")
B("**Comprensión de los datos ampliada:** además de los conteos, se midió el contenido de las imágenes. Hay dos modos de archivo (gris y RGB) que se reparten distinto por clase, el brillo y el encuadre cambian con la clase, el 8,3 % de las máscaras tiene más de una región y la prueba oficial no es una copia exacta del entrenamiento. Todo esto se tradujo en decisiones y riesgos para el Sprint 2.")
B("**Definición del proyecto:** Acta de Constitución, Product Backlog de 17 historias en Trello con criterios de aceptación, cronograma (Gantt) "
  "y las plantillas de gestión del repositorio del docente llenas.")
B("**Exploración de BRISC 2025 con los datos reales:** 6.000 imágenes (5.000 de entrenamiento y 1.000 de prueba) y 4.793 pares imagen–máscara. "
  "Los conteos por clase y partición coinciden con el artículo del dataset (diferencia 0 en las 8 combinaciones).")
B("**Hallazgos que cambiaron el código:** las máscaras no son estrictamente binarias, las imágenes tienen 183 tamaños distintos y las 1.207 "
  "imágenes sin tumor no tienen máscara.")
B("**Partición reproducible:** 3.999 imágenes de entrenamiento, 1.001 de validación (20 % en cada clase y plano de corte) y 1.000 de prueba oficial intacta, "
  "agrupando imágenes casi idénticas para reducir la fuga entre entrenamiento y validación.")
B("**Cargador de datos** probado con un lote real, y todo el flujo reproducible con dos comandos (`descargar_brisc.py` y `ejecutar_sprint1.py`).")
B("**Riesgo detectado en el dataset:** entre 113 y 123 de las 1.000 imágenes de la prueba oficial tienen un vecino casi idéntico en el entrenamiento, "
  "por lo que el resultado final podría salir algo optimista y debe reportarse con esa advertencia.")
NOTE("**Aclaración importante:** en este sprint **no se entrenó ningún modelo** y no hay métricas de desempeño (accuracy, recall, IoU). "
     "Todos los resultados son de análisis y preparación de datos. El entrenamiento corresponde al Sprint 2.")

# ================================================================ 2. CONTEXTO
H1("2. Comprensión del negocio")
P("Esta es la primera tarea de CRISP-ML(Q): antes de mirar los datos hay que dejar claro **qué problema se quiere resolver, para quién, con qué recursos y cómo se medirá el éxito**. "
  "Este capítulo recoge esa comprensión. Se apoya en cifras públicas con su fuente y distingue con claridad lo que es un **dato verificado**, lo que es una **decisión del equipo** y lo que es un **supuesto ilustrativo**.")
NOTE("**Sobre los \"clientes\":** el proyecto es académico y **no tiene un cliente real, patrocinador ni Dueño del Producto**; el equipo define el backlog y el docente solo revisa. "
     "Por eso no existe ninguna clínica con un número real de clientes o pacientes. Para describir a quién iría dirigida una versión real del sistema se usa un **escenario de referencia hipotético** "
     "(sección 2.4), cuyos números son supuestos explícitos del equipo y no datos de ninguna institución.")

H2("2.1 El problema y su contexto")
P("Un tumor cerebral se detecta y se estudia principalmente con **resonancia magnética (MRI)**. Quien interpreta el estudio, normalmente un médico radiólogo, debe decidir si hay un tumor, "
  "de qué tipo es y dónde está. Esa información orienta los pasos siguientes (por ejemplo, la derivación a neurocirugía u oncología). "
  "Los tres tipos de tumor que cubre BRISC son:")
T([["Tipo", "Qué es (resumen)", "Por qué importa distinguirlo"],
   ["Glioma", "Tumor que se origina en las células de soporte del cerebro (células gliales).", "Suele requerir un manejo distinto al de otros tumores y su extensión es clave para planificar el tratamiento."],
   ["Meningioma", "Tumor que se origina en las meninges, las membranas que recubren el cerebro.", "Es frecuente y muchas veces de crecimiento lento; su ubicación y tamaño orientan el seguimiento o la cirugía."],
   ["Tumor de hipófisis", "Tumor de la glándula hipófisis, en la base del cerebro.", "Puede afectar hormonas y visión; se localiza en una región pequeña y específica."]],
  [3.2, 6.6, 6.8], size=8.5)
P("Estas descripciones son generales y de carácter divulgativo; el proyecto **no** pretende reemplazar la interpretación médica.", italic=True, size=9.5, color=GRAY)
H3("Datos públicos que dimensionan el problema (verificados)")
T([["Dato", "Valor", "Fuente"],
   ["Población de Perú (estimación 2024)", "34.217.853 personas", "GLOBOCAN 2024, ficha de Perú (IARC / OMS)"],
   ["Casos nuevos de cáncer, todos los tipos (2024)", "77.485", "GLOBOCAN 2024, ficha de Perú"],
   ["Casos nuevos de cáncer de cerebro y sistema nervioso central (2024)", "1.616 (puesto 14 entre los cánceres; 2,1 % del total)", "GLOBOCAN 2024, ficha de Perú"],
   ["Muertes por cáncer de cerebro y sistema nervioso central (2024)", "1.147 (puesto 11; 3,1 % de las muertes por cáncer)", "GLOBOCAN 2024, ficha de Perú"],
   ["Casos prevalentes a 5 años, cerebro y sistema nervioso central", "5.545 (16,2 por 100.000 habitantes)", "GLOBOCAN 2024, ficha de Perú"],
   ["Incidencia estimada de cáncer de cerebro y SNC en Perú (calculada)", "4,7 por 100.000 habitantes al año (1.616 ÷ 34.217.853)", "Cálculo propio con las dos cifras anteriores de GLOBOCAN 2024"],
   ["Referencia internacional: todos los tumores primarios de cerebro y SNC, incluidos los no malignos (EE. UU., 2014–2018)", "≈ 24 por 100.000 al año (24,25 ajustada por edad); los meningiomas son ≈ 39 % de los tumores primarios (informe 2015–2019)", "Informes CBTRUS; cifras tomadas de extractos de búsqueda, **no se abrió el informe original**"],
   ["Médicos radiólogos en Perú", "495", "Observatorio de Recursos Humanos en Salud del Minsa, citado por Panamericana (nov. 2025)"],
   ["Tecnólogos médicos y médicos patólogos en Perú", "898 y 154", "Observatorio de Recursos Humanos en Salud del Minsa, citado por Panamericana (nov. 2025)"],
   ["Concentración de radiólogos y tecnólogos médicos en Lima", "54,8 %", "Observatorio de Recursos Humanos en Salud del Minsa, citado por Panamericana (nov. 2025)"]],
  [6.4, 4.8, 5.4], size=8.5)
P("**Cómo leer estas cifras (con cuidado):**", keep=True)
B("Son **estimaciones de GLOBOCAN** (modelos del IARC), no un conteo de pacientes de una clínica. Cubren el cáncer de cerebro y sistema nervioso central; "
  "no necesariamente incluyen los tumores benignos (como muchos meningiomas o tumores de hipófisis), así que **no equivalen al número de pacientes con las clases de BRISC**. No se usan para calcular una demanda.")
B("Los datos de personal médico vienen de una nota periodística que cita al Observatorio del Minsa; no se contrastaron con el documento original del Minsa.")
B("Lo que sí muestran en conjunto: el número de especialistas que interpretan imágenes es pequeño frente a la población y está concentrado en Lima. "
  "Esa es la motivación general del proyecto; el proyecto **no mide ni demuestra** una reducción de esa carga.")

H2("2.2 El sector privado de salud: establecimientos, pacientes y resonancia (datos públicos)")
P("Para saber a quién iría dirigido el sistema hay que mirar dónde se hacen e interpretan las resonancias. Estas son las cifras públicas que se pudieron leer y verificar, con su año y su límite. "
  "Donde no existe una cifra pública, se dice y no se inventa.")
H3("Establecimientos y aseguramiento (SuSalud, mayo de 2022)")
T([["Dato", "Valor", "Observación"],
   ["Establecimientos de salud activos y categorizados", "19.067 (97 % de primer nivel; 3 % de segundo y tercer nivel)", "Registro Nacional de IPRESS, mayo de 2022"],
   ["Según sector", "Público 8.866 (46 %) · EsSalud 381 (2 %) · privado 9.820 (52 %)", "Suman 19.067"],
   ["Privados de primer nivel", "9.508, en su mayoría consultorios médicos y de otros profesionales", "No son centros de diagnóstico por imágenes"],
   ["Privados de segundo y tercer nivel (clínicas y hospitales privados)", "≈ 312", "**Cálculo propio:** 9.820 − 9.508. Es el universo donde podría haber resonadores; no todos los tienen"],
   ["Establecimientos activos sin categoría", "5.179 (95 % privados)", "Mayormente consultorios, sin clasificar"],
   ["Camas hospitalarias", "1,1 por 1.000 habitantes (promedio de América Latina y el Caribe: 1,9)", "El sector público y EsSalud concentran el 84 % de las camas"],
   ["Afiliados al SIS", "≈ 24,7 millones (69,1 % de la población)", "Una persona puede tener más de un seguro"],
   ["Afiliados a EsSalud", "28,5 % de la población", "Seguro social de trabajadores formales"],
   ["Población con aseguramiento privado (sin contar EPS)", "6,7 % a nivel nacional; 14,5 % en Lima Metropolitana; 9,4 % en Arequipa", "Es el segmento que se atiende en clínicas privadas"]],
  [5.2, 5.8, 5.6], size=8.5)
P("Fuente: Videnza Consultores para ComexPerú, Situación actual del mercado de salud peruano (agosto de 2022), con datos de SuSalud y del Minsa. "
  "Son cifras de 2022: no se actualizaron a 2026.", italic=True, size=9, color=GRAY)
FIG(FIGS + "sector_privado.png", 15.0, "Establecimientos de salud por sector y población por tipo de seguro (Perú, 2022).")
H3("Resonancia magnética")
T([["Dato", "Valor", "Observación"],
   ["Atenciones con tomógrafos y resonadores del Minsa en el primer semestre", "31.101 (25.346 de tomografía y 5.755 resonancias magnéticas)", "Nota de Panamericana; el año no figura en el texto y la nota se publicó hacia agosto de 2025"],
   ["Resonancias por hospital en ese semestre", "Cayetano Heredia 5.218 · María Auxiliadora 517", "Suman 5.735, no 5.755: la nota no reconcilia del todo. Los equipos son de 1,5 Tesla"],
   ["Carga mensual de un equipo público muy usado", "≈ 870 resonancias al mes", "**Cálculo propio:** 5.218 ÷ 6 meses (Hospital Cayetano Heredia)"],
   ["Médicos radiólogos / tecnólogos médicos / patólogos", "495 / 898 / 154", "Observatorio de RR. HH. en Salud del Minsa, citado por Panamericana (nov. 2025)"]],
  [5.2, 5.8, 5.6], size=8.5)
H3("Lo que no se encontró (y no se inventó)")
B("Número total de resonadores en el sector privado, y cuántos estudios de resonancia cerebral se hacen al año en Perú.")
B("Cuántos pacientes atiende en promedio una clínica privada, y cuántas resonancias hace.")
B("Los registros que podrían tener esos datos son el Registro Nacional de IPRESS (RENIPRESS) de SuSalud y los datos abiertos del Minsa y de SuSalud; consultarlos queda como mejora si el proyecto pasara a un producto real.")

H2("2.3 A quién está dirigido")
P("Conviene separar quién usa realmente el sistema hoy de quién podría usarlo si el proyecto se convirtiera en un producto:", keep=True)
T([["Grupo", "Relación con el proyecto", "Qué obtiene o qué necesita"],
   ["**Equipo del proyecto** (4 integrantes)", "Usuario real de la demostración y responsable de construirla", "Aprender y demostrar el ciclo completo de un sistema de IA: datos, modelos, evaluación y despliegue."],
   ["**Docente revisor** (Ing. Hugo Espetia Huamanga)", "Evalúa el proyecto en cada sprint y en la entrega final; no es patrocinador ni Dueño del Producto", "Evidencia del proceso (Scrum + CRISP-ML(Q)), resultados medidos con rigor y una demostración funcional."],
   ["**Radiólogos y tecnólogos médicos** (hipotético)", "Usuarios potenciales de una versión futura, **no consultados** en este proyecto", "Una segunda opinión visual: clase probable del tumor y su contorno sobre la imagen, siempre revisable."],
   ["**Neurólogos, neurocirujanos y oncólogos** (hipotético)", "Reciben el informe del radiólogo; **no consultados**", "Información clara y rápida sobre el tipo y la ubicación del tumor."],
   ["**Centros de diagnóstico por imágenes y clínicas privadas** (hipotético)", "Quien compraría o implementaría una versión real; **no consultados**", "Ver el escenario de referencia de la sección 2.4."],
   ["**Pacientes** (hipotético)", "Beneficiarios indirectos de una versión futura", "Que su estudio se interprete con calidad y en un tiempo razonable. No usan la demostración."]],
  [4.4, 5.6, 6.6], size=8.5)

H2("2.4 Escenario de referencia: clínica privada con resonancia (modelo con supuestos)")
P("Como el proyecto no tiene una clínica real, se arma un **modelo** de una clínica privada genérica con resonancia, en tres tamaños. "
  "**Es un modelo del equipo con supuestos explícitos; ninguna cifra de esta sección es un dato medido de una institución.** "
  "Cada supuesto indica si se apoya en un dato público; si se consiguiera una clínica, se reemplazan por los valores reales.")
T([["Parámetro", "Pequeña", "Mediana", "Grande", "Base del supuesto"],
   ["Equipos de resonancia", "1", "2", "4", "Supuesto"],
   ["Estudios de MRI al mes (todos los tipos)", "400", "1.200", "3.000", "Supuesto. Referencia: un equipo público muy usado hace ≈ 870 al mes (sección 2.2)"],
   ["Parte de los estudios que son de cerebro", "25 %", "25 %", "25 %", "Supuesto; sin dato público"],
   ["**Estudios de MRI cerebral al mes**", "**100**", "**300**", "**750**", "Cálculo: estudios × 25 %"],
   ["**Estudios de MRI cerebral al año**", "**1.200**", "**3.600**", "**9.000**", "Cálculo: × 12. Se asume un estudio por paciente"],
   ["Tiempo de lectura e informe por estudio cerebral", "20 a 30 min", "20 a 30 min", "20 a 30 min", "Supuesto"],
   ["**Horas médicas de lectura cerebral al mes**", "**33 a 50**", "**100 a 150**", "**250 a 375**", "Cálculo: estudios × tiempo"],
   ["Radiólogos que informan", "1 a 2", "2 a 3", "4 a 6", "Supuesto"]],
  [4.6, 1.9, 1.9, 1.9, 6.3], size=8)
FIG(FIGS + "escenario_clinica.png", 15.0, "Escenario de clínica privada con resonancia: resonancias cerebrales al año y horas de lectura al mes (modelo con supuestos).")
P("**Quiénes son los \"clientes\" en este modelo.** Son los pacientes que llegan con indicación de resonancia cerebral: entre 1.200 y 9.000 al año según el tamaño, casi todos asegurados privados o con seguro "
  "que deriva a clínicas (en 2022, el 6,7 % de la población tenía aseguramiento privado sin contar EPS, y el 14,5 % en Lima Metropolitana). "
  "El comprador de una versión real no sería el paciente sino la clínica, entre las ≈ 312 clínicas y hospitales privados de segundo y tercer nivel del país en 2022, de las cuales no se sabe cuántas tienen resonador.")
H3("Cuántos tumores esperaría ver esa clínica")
P("La incidencia es la base. Con la cifra peruana de GLOBOCAN (4,7 casos nuevos de cáncer de cerebro y SNC por 100.000 habitantes al año) y la referencia internacional que incluye tumores no malignos (≈ 24 por 100.000), "
  "una clínica que atienda a **100.000 personas** vería, en un año, del orden de **5 a 24 casos nuevos**. Es un dato ilustrativo: el número de personas atendidas es un supuesto, y las cifras de incidencia no son de esa clínica.")
P("**Conclusión de negocio:** los tumores son raros frente al volumen de estudios cerebrales (miles al año frente a decenas de casos nuevos). "
  "Casi todas las resonancias cerebrales de esa clínica serían normales o con otros hallazgos, así que un sistema de apoyo útil debe equivocarse poco en las imágenes sin tumor.")
H3("Por qué importa: la prevalencia de BRISC no es la de una clínica")
P("En BRISC el 80 % de las imágenes tiene tumor (4.793 de 6.000), porque es un dataset armado para entrenar. En una clínica real los tumores serían una fracción pequeña de los estudios. "
  "Cuando algo es raro, aun un modelo con buenas cifras genera muchas falsas alarmas. Ejemplo ilustrativo (los porcentajes son hipotéticos, no metas del proyecto):", keep=True)
T([["Supuesto", "Valor"],
   ["Prevalencia de tumor entre los estudios cerebrales", "3 %"],
   ["Sensibilidad del modelo (detecta tumor cuando lo hay)", "95 %"],
   ["Especificidad del modelo (dice \"sin tumor\" cuando no lo hay)", "95 %"],
   ["**De cada 100 estudios señalados como tumor, los que realmente lo tienen**", "**≈ 37**"]],
  [11.6, 5.0], size=9, align_right=(1,))
P("Cálculo: 0,95 × 0,03 = 0,0285 verdaderos positivos y 0,05 × 0,97 = 0,0485 falsos positivos por estudio; 0,0285 ÷ (0,0285 + 0,0485) ≈ 37 %. "
  "Por eso las metas del proyecto (accuracy y recall por clase, medidas en un dataset con mayoría de tumores) **no dicen cómo se comportaría el sistema en una clínica**. "
  "Para una demostración académica es suficiente; para un uso real habría que medir precisión y especificidad con la prevalencia real del lugar.")
FIG(FIGS + "prevalencia_ppv.png", 11.5, "Cuántos de los estudios señalados como tumor lo tienen de verdad según la prevalencia (valores hipotéticos de sensibilidad y especificidad).")
P("**Dónde encajaría el sistema en el flujo de esa clínica (si se llegara a un producto):**", keep=True)
T([["Paso", "Qué ocurre", "Qué haría el sistema"],
   ["1", "El paciente se realiza la resonancia y la imagen se guarda en el archivo de imágenes de la clínica.", "—"],
   ["2", "El radiólogo abre el estudio.", "Un módulo ofrecería, como apoyo, la clase probable (glioma, meningioma, hipófisis o sin tumor), su probabilidad y la máscara del tumor."],
   ["3", "El radiólogo interpreta la imagen y **decide**.", "No decide ni informa: solo muestra una sugerencia que el médico puede aceptar o descartar."],
   ["4", "Se emite el informe al médico tratante.", "—"]],
  [1.4, 7.2, 8.0], size=8.5)
P("**Lo que el proyecto sí entrega** es solo una demostración académica de los pasos 2 y 3 sobre imágenes de un dataset público; no se integra con equipos reales, con sistemas clínicos ni con el formato de imagen médica de las clínicas.")

H2("2.5 Qué se construye y para qué")
B("**¿Para qué?** Mostrar, con fines académicos, cómo una red neuronal puede apoyar la interpretación de resonancias magnéticas (MRI) cerebrales: "
  "identifica el tipo de tumor (glioma, meningioma, tumor de hipófisis o sin tumor) y delimita su contorno.")
B("**¿Para dónde?** Una aplicación web de demostración en la que quien la usa (el equipo y el docente revisor) sube una MRI y recibe la clase, "
  "la probabilidad y la máscara del tumor. Uso académico y de investigación con imágenes T1 con contraste; no está destinado a uso clínico.")
B("**Sin patrocinador ni Dueño del Producto:** es un proyecto académico; los 4 integrantes definen y priorizan el backlog y el docente solo revisa.")

H2("2.6 Objetivos de negocio, objetivos de aprendizaje automático y criterios de éxito")
T([["Nivel", "Objetivo", "Criterio de éxito"],
   ["Del proyecto (académico)", "Completar un sistema de IA de punta a punta aplicando Scrum + CRISP-ML(Q), con evidencia del proceso en cada sprint.", "Tres presentaciones al docente (12, 23 y 30 de octubre) y entrega antes del 2 de noviembre de 2026."],
   ["Del producto (demostración)", "Una aplicación web que, dada una MRI, muestre clase, probabilidad y máscara del tumor.", "Demostración funcional en la presentación final, con modelos cargados desde archivos exportados."],
   ["De aprendizaje automático: clasificación", "Clasificar la imagen en 4 clases.", "Accuracy ≥ 95 % y recall ≥ 90 % en cada clase, sobre la prueba oficial de BRISC."],
   ["De aprendizaje automático: segmentación", "Delimitar el tumor en las imágenes con tumor.", "IoU ≥ 70 % sobre la prueba oficial, solo imágenes con tumor."]],
  [3.6, 7.0, 6.0], size=8.5)
P("Las metas numéricas son **propuestas del equipo** y se fijaron tomando como referencia los resultados publicados en el artículo del dataset (EfficientNetB0 99,2 % y ResNet50 98,2 % de accuracy; "
  "U-Net con mIoU ponderado de 75,7 %). Esos resultados **no son propios**. Las metas se miden **solo sobre la prueba oficial** y todavía no se han evaluado: en el Sprint 1 no se entrenó ningún modelo.")

H2("2.7 Evaluación de la situación")
T([["Aspecto", "Situación del proyecto"],
   ["Recursos humanos", "4 estudiantes de Ingeniería de Sistemas – Inteligencia Artificial, con dedicación parcial. Roles: Gerente de Proyecto, Scrum Master y dos integrantes de desarrollo (sección 2.11). El docente revisa."],
   ["Datos", "Dataset público BRISC 2025: 6.000 imágenes T1 con contraste, 4.793 pares imagen–máscara. Licencia CC BY 4.0. Sin ID de paciente."],
   ["Cómputo", "Equipo personal con GPU NVIDIA RTX 4050 de 6 GB (PyTorch instalado solo para CPU; hay que instalar la versión CUDA). Google Colab gratuito como respaldo."],
   ["Software", "Python 3.13, bibliotecas de código abierto, GitHub, Trello y herramientas gratuitas."],
   ["Costos", "Costo programado de S/ 0: dataset, cómputo y alojamiento con planes gratuitos."],
   ["Plazo", "33 días (1 de octubre al 2 de noviembre de 2026), tres sprints y un margen de un día (31 de octubre)."],
   ["Requisitos y restricciones", "Metodología híbrida Scrum + CRISP-ML(Q); entregables del docente (acta, backlog, informes, presentaciones); uso solo académico; solo imágenes T1 con contraste."],
   ["Supuestos", "El dataset público es representativo para una demostración; las arquitecturas probadas en el artículo se pueden reproducir en una GPU de 6 GB; el equipo puede dedicar tiempo a los sprints según el cronograma."],
   ["Riesgos principales", "Fuga entre pacientes, cambio de dominio, arquitecturas inestables, GPU limitada, retraso del cronograma y uso indebido de la demo (detalle en la sección 8)."],
   ["Beneficios esperados", "Aprendizaje del ciclo completo, una demostración funcional y un registro reproducible del proceso. No se estiman beneficios clínicos ni económicos."]],
  [3.6, 13.0], size=8.5)

H2("2.8 Factibilidad")
P("CRISP-ML(Q) pide comprobar, antes de modelar, que el proyecto es factible. La conclusión del equipo, con la evidencia disponible al cierre del Sprint 1, es que **es factible con el alcance acotado**:", keep=True)
T([["Pregunta", "Evidencia", "Valoración"],
   ["¿Hay datos suficientes y de calidad?", "6.000 imágenes y 4.793 pares; conteos iguales a los del artículo; 0 archivos corruptos; problemas conocidos y tratados (máscaras no binarias, tamaños variables, duplicados).", "Sí, con limitaciones declaradas"],
   ["¿Es el problema abordable con las técnicas previstas?", "El artículo de BRISC reporta alta exactitud con CNN preentrenadas y un mIoU de 75,7 % con U-Net sobre este mismo dataset.", "Sí (referencia externa; falta probarlo)"],
   ["¿Hay cómputo?", "GPU local de 6 GB; modelos medianos con transfer learning caben; Colab de respaldo.", "Sí, con riesgo de memoria"],
   ["¿Alcanza el tiempo?", "Dos sprints restantes con cronograma ajustado y un día de margen.", "Ajustado; el modelo multitarea queda como opcional"],
   ["¿Es aceptable el riesgo de uso?", "Uso solo académico, aviso visible de que no es una herramienta de diagnóstico.", "Sí, con aviso en la aplicación"]],
  [4.6, 8.2, 3.8], size=8.5)

H2("2.9 Consideraciones éticas y de uso")
B("**No es un dispositivo médico.** El sistema no diagnostica ni debe usarse para decisiones clínicas; la aplicación mostrará un aviso visible de demostración académica.")
B("**Datos públicos:** se usa un dataset de acceso abierto con licencia CC BY 4.0; hay que citar a los autores de BRISC y de las fuentes originales.")
B("**Sesgos y límites de los datos:** solo hay imágenes T1 con contraste y no se conoce el detalle de pacientes ni de equipos; los resultados pueden no generalizar a otras secuencias, otros centros o poblaciones.")
B("**Honestidad en los resultados:** la prueba oficial se usa una sola vez al final, y se informará que parte de ella se parece a imágenes del entrenamiento (sección 4.3).")

H2("2.10 Metodología híbrida: Scrum + CRISP-ML(Q)")
P("Scrum organiza el trabajo del equipo (backlog, sprints, revisión con el docente y retrospectiva). CRISP-ML(Q) organiza el ciclo de vida del modelo; "
  "la \"Q\" de calidad se aplica con criterios de aceptación en cada historia y con el registro de riesgos. Cada sprint cubre fases de CRISP-ML(Q):")
T([["Sprint", "Trabajo", "Presentación", "Fase(s) CRISP-ML(Q)"],
   ["1", "1 – 8 oct 2026", "lun 12 oct", "Comprensión del negocio y de los datos + preparación de los datos"],
   ["2", "10 – 22 oct 2026", "vie 23 oct", "Modelado"],
   ["3", "24 – 29 oct 2026", "vie 30 oct", "Evaluación + despliegue (MLOps básico)"]],
  [1.6, 3.6, 3.0, 8.4])
P("Margen de seguridad: 31 de octubre. Fecha límite del docente: lunes 2 de noviembre de 2026. La fase de monitoreo y mantenimiento se cubre solo en lo básico por el plazo del proyecto (33 días).")

H2("2.11 Equipo y roles")
T([["Integrante", "Rol"],
   ["Coavoy Cruz Joseph Gabriel", "Gerente de Proyecto"],
   ["Cuchuyrrumi Mamani Manuel Rodrigo", "Scrum Master"],
   ["Huallpatuiro Rafaile Brayan", "Equipo de desarrollo"],
   ["Mamani Acuña Frank Joseph", "Equipo de desarrollo"]],
  [8.0, 8.6])

# ================================================================ 3. PLAN
H1("3. Plan del Sprint 1")
P("**Meta del sprint:** entender el problema y el dataset BRISC a fondo, y dejar los datos listos y particionados sin fuga evidente.")
T([["Historia", "Fase CRISP-ML(Q)", "Gantt", "Puntos", "Prioridad", "Estado"],
   ["HU01 · Definir el problema y el alcance", "Comprensión del negocio y de los datos", "1 oct", "2", "Alta", "Terminada"],
   ["HU02 · Descargar y explorar BRISC (EDA)", "Comprensión del negocio y de los datos", "2 – 5 oct", "5", "Alta", "Terminada"],
   ["HU03 · Preparar los datos y definir la partición", "Preparación de los datos", "6 – 8 oct", "5", "Alta", "Terminada"],
   ["HU04 · Presentación del Sprint 1", "Revisión del sprint (Scrum)", "12 oct", "2", "Media", "Terminada"],
   ["**Total**", "", "", "**14**", "", ""]],
  [4.7, 4.2, 1.7, 1.6, 1.7, 2.7], align_right=(3,))
P("La estimación de esfuerzo del Backlog del Sprint (19 tareas) suma **53 horas**; es una estimación propuesta por el equipo en la planificación, no un registro de horas trabajadas. "
  "Las tareas por historia fueron:", keep=True)
T([["Tarea", "Historia", "Descripción", "Est. (h)"],
   ["1.1", "HU01", "Redactar y revisar el Acta de Constitución (título, ¿para qué?, ¿para dónde?, riesgos, hitos)", "6"],
   ["1.2", "HU01", "Fijar las metas medibles (accuracy, recall por clase e IoU) con referencia del artículo de BRISC", "2"],
   ["1.3", "HU01", "Armar el Product Backlog en Trello y el cronograma (Gantt)", "4"],
   ["2.1", "HU02", "Descargar BRISC en el equipo y preparar el notebook", "2"],
   ["2.2", "HU02", "Verificar la estructura de carpetas y leer manifest.csv", "2"],
   ["2.3", "HU02", "Contar imágenes por clase, partición y plano, y compararlas con el artículo", "3"],
   ["2.4", "HU02", "Medir dimensiones de imagen y porcentaje del tumor en cada máscara", "3"],
   ["2.5", "HU02", "Galería de imagen con máscara por clase", "2"],
   ["2.6", "HU02", "Confirmar si las imágenes sin tumor tienen máscara y la licencia del dataset", "1"],
   ["2.7", "HU02", "Resumir los hallazgos del análisis exploratorio", "2"],
   ["3.1", "HU03", "Verificar archivos corruptos y pares imagen-máscara", "3"],
   ["3.2", "HU03", "Agrupar imágenes muy parecidas (hash perceptual) para reducir la fuga entre particiones", "4"],
   ["3.3", "HU03", "Definir la partición de entrenamiento y validación (estratificada por clase y plano, semilla fija)", "3"],
   ["3.4", "HU03", "Guardar las listas de la partición y versionarlas en GitHub", "1"],
   ["3.5", "HU03", "Escribir el cargador de datos y probar que carga un lote", "4"],
   ["3.6", "HU03", "Fijar el preprocesamiento (tamaño de entrada y normalización)", "2"],
   ["4.1", "HU04", "Armar las diapositivas de la revisión del Sprint 1", "5"],
   ["4.2", "HU04", "Hacer un ensayo de la presentación", "2"],
   ["4.3", "HU04", "Hacer la retrospectiva del Sprint 1", "2"]],
  [1.3, 1.6, 12.2, 1.5], size=8.5, align_right=(3,))

# ================================================================ 4. TRABAJO
H1("4. Trabajo realizado")
H2("4.1 HU01 · Definir el problema y el alcance")
P("Entregables de la historia:", keep=True)
B("**Acta de Constitución** (firmada): título del proyecto con su ¿para qué? y ¿para dónde?, propósito, alcance (qué incluye y qué no), restricciones, "
  "metas medibles, riesgos y oportunidades, hitos, presupuesto y equipo. No hay patrocinador ni Dueño del Producto; el docente es el revisor.")
B("**Product Backlog en Trello** con 17 historias (HU01 a HU17) organizadas en 4 épicas: EP1 Entender el negocio y los datos, EP2 Modelar, "
  "EP3 Evaluación y despliegue (MLOps) y EP4 Informe y presentación final. Cada tarjeta trae historia de usuario, prioridad, estimación en puntos, "
  "dependencias, fase CRISP-ML(Q) y un checklist de criterios de aceptación. Además hay tarjetas de proyecto y metodología, y una de meta del sprint.")
B("**Cronograma:** diagrama de Gantt y archivo CSV para la herramienta de cronogramas, con los tres sprints y sus presentaciones.")
B("**Plantillas de gestión llenas** (formato del repositorio del docente): Backlog del Producto, Registro de Interesados, Planificación de Lanzamiento, "
  "Backlog Detallado, Backlog del Sprint, Cronograma de hitos, Presupuesto (costo programado de S/ 0: dataset, cómputo y alojamiento con planes gratuitos), "
  "Registro de Riesgos, Informe de Estado, Minuta, Registro de Impedimentos y las presentaciones de revisión y retrospectiva.")
B("**Revisión general del Acta contra el artículo de BRISC:** se corrigieron la licencia, las referencias de resultados y el presupuesto, "
  "para que no queden datos inventados ni contradicciones.")

H2("4.2 HU02 · Descargar y explorar BRISC")
H3("Qué es BRISC 2025")
P("BRISC 2025 es un dataset público de resonancias magnéticas cerebrales **T1 con contraste**: 6.000 cortes en tres planos (axial, coronal y sagital) y cuatro clases "
  "(glioma, meningioma, tumor de hipófisis y sin tumor). Incluye una tarea de clasificación y una de segmentación, con una partición oficial de 5.000 imágenes de "
  "entrenamiento y 1.000 de prueba. Según el artículo, deriva de datasets públicos previos (Nickparvar, Cheng, SARTAJ y Br35H). "
  "**Licencia: CC BY 4.0** (metadatos de Kaggle), por lo que hay que citar a los autores. Los nombres de archivo siguen el patrón "
  "`brisc2025_<partición>_<índice>_<código>_<plano>_<secuencia>.jpg`.")
H3("Descarga y estructura")
P("El dataset se descargó con el script `descargar_brisc.py` a una carpeta fuera de OneDrive (`~/datasets/brisc2025`), de unos 277 MB y 15.586 archivos de imagen. "
  "El notebook verificó la estructura: carpetas `classification_task` y `segmentation_task` más el archivo `manifest.csv` (con su suma SHA-256 por archivo).")
T([["Carpeta", "Entrenamiento", "Prueba"],
   ["classification_task / glioma", "1.147", "254"],
   ["classification_task / meningioma", "1.329", "306"],
   ["classification_task / pituitary (hipófisis)", "1.457", "300"],
   ["classification_task / no_tumor (sin tumor)", "1.067", "140"],
   ["segmentation_task / images", "3.933", "860"],
   ["segmentation_task / masks", "3.933", "860"]],
  [8.6, 4.0, 4.0], align_right=(1, 2))
H3("Conteos por clase y plano de corte")
P("Los conteos por clase y partición son **idénticos a los del artículo** (diferencia 0 en las 8 combinaciones):", keep=True)
T([["Clase", "Entrenamiento", "Prueba", "Total", "Axial", "Coronal", "Sagital"],
   ["Glioma", "1.147", "254", "1.401", "479", "511", "411"],
   ["Meningioma", "1.329", "306", "1.635", "560", "512", "563"],
   ["Tumor de hipófisis", "1.457", "300", "1.757", "550", "600", "607"],
   ["Sin tumor", "1.067", "140", "1.207", "404", "358", "445"],
   ["**Total**", "**5.000**", "**1.000**", "**6.000**", "**1.993**", "**1.981**", "**2.026**"]],
  [4.2, 2.5, 1.8, 1.8, 2.0, 2.1, 2.2], align_right=(1, 2, 3, 4, 5, 6))
P("Los conteos por **plano de corte** difieren levemente de la Tabla 2 del artículo. Se usan los valores observados:", keep=True)
T([["Partición / plano", "Artículo", "Observado", "Diferencia"],
   ["Entrenamiento · axial", "1.550", "1.595", "+45"],
   ["Entrenamiento · coronal", "1.660", "1.676", "+16"],
   ["Entrenamiento · sagital", "1.790", "1.729", "−61"],
   ["Prueba · axial", "387", "398", "+11"],
   ["Prueba · coronal", "316", "305", "−11"],
   ["Prueba · sagital", "297", "297", "0"]],
  [6.0, 3.4, 3.6, 3.6], align_right=(1, 2, 3))
P("El código de clase en el nombre del archivo es `gl` (glioma), `me` (meningioma), `pi` (hipófisis) y `no` (sin tumor); el README del dataset menciona `nt` para sin tumor, "
  "pero los archivos reales usan `no`. Cada código corresponde a una sola clase. La composición está levemente desbalanceada (de 1.207 a 1.757 imágenes por clase).")
FIG(SAL + "figuras_eda/conteo_clase_particion.png", 11.5, "Figura 1. Imágenes por clase y partición oficial.")
FIG(SAL + "figuras_eda/planos_por_clase.png", 11.5, "Figura 2. Planos de corte por clase.")

H3("Segmentación: pares imagen–máscara")
B("Hay **4.793 pares** imagen–máscara completos (3.933 de entrenamiento y 860 de prueba), igual que el artículo. No hay máscaras sin imagen ni imágenes sin máscara dentro de la tarea de segmentación.")
B("Distribución de los pares por clase: glioma 1.401, meningioma 1.635 y tumor de hipófisis 1.757.")
B("**Las 1.207 imágenes sin tumor no tienen máscara** (0 de 1.207). La página de Kaggle habla de 6.000 imágenes y máscaras, pero los archivos confirman lo que decía el artículo. "
  "Por eso la segmentación se entrena y se evalúa solo con las 4.793 imágenes con tumor.")
B("Ninguna máscara está vacía.")
H3("Formato de las máscaras y tamaño de las imágenes")
B("**Máscaras no binarias:** las 4.793 máscaras (archivos PNG) contienen valores intermedios (normalmente de 0 a 6 y, en algunos casos, valores altos como 250), "
  "probablemente por remuestreo. No es un error de los datos, pero obliga a binarizar: se consideran tumor los píxeles > 127.")
B("**Tamaños variables:** hay 183 combinaciones de ancho × alto. La más común es 512×512 (4.881 imágenes); siguen 369×369 (363) y 216×369 (354). "
  "Las máscaras son de 512×512 en 4.504 casos, 256×256 en 30 y 225×225 en 15 (entre las tres más frecuentes). No hay una relación de aspecto única.")
H3("Tamaño del tumor")
P("El tumor ocupa una fracción muy pequeña de la imagen (en mediana, entre el 0,9 % y el 1,7 %), de modo que en la segmentación el fondo domina ampliamente. "
  "Esto condiciona la elección de la función de pérdida en el Sprint 2.", keep=True)
T([["Clase", "Pares", "Media", "Mediana", "Mínimo", "Máximo"],
   ["Glioma", "1.401", "1,87 %", "1,52 %", "0,05 %", "9,77 %"],
   ["Meningioma", "1.635", "2,25 %", "1,74 %", "0,01 %", "17,35 %"],
   ["Tumor de hipófisis", "1.757", "1,04 %", "0,86 %", "0,06 %", "6,49 %"]],
  [4.6, 2.2, 2.4, 2.4, 2.4, 2.6], align_right=(1, 2, 3, 4, 5))
FIG(SAL + "figuras_eda/tamano_tumor.png", 11.5, "Figura 3. Distribución del tamaño del tumor (fracción de píxeles de la imagen) por clase.")
FIG(SAL + "figuras_eda/galeria_imagen_mascara.png", 10.2, "Figura 4. Imagen y máscara (en rojo) por clase y plano; la fila inferior es la clase sin tumor, que no tiene máscara.")

H3("manifest.csv y duplicados")
P("El manifest tiene 16 columnas (ruta, nombre, tarea, partición, clase, plano, secuencia, indicador de máscara, imagen vinculada, ancho, alto, tamaño y SHA-256). "
  "Contiene 6.000 imágenes de clasificación; su tarea de segmentación repite las mismas imágenes, por lo que un conteo de duplicados sobre todo el archivo es engañoso "
  "(la primera versión del notebook reportaba 9.639 hashes repetidos por ese motivo). Corregido el filtro, **96 archivos de clasificación tienen un SHA-256 repetido** "
  "(duplicados exactos) y **7 de esos hashes aparecen tanto en entrenamiento como en prueba**.")
H3("Limitaciones registradas")
B("El dataset no trae ID de paciente: cortes de un mismo paciente pueden estar en particiones distintas.")
B("Solo hay imágenes T1 con contraste: los modelos pueden rendir peor con otras secuencias (T2, sin contraste) o con otros equipos.")
B("La clase sin tumor incluye cerebros sanos y lesiones no neoplásicas (según el artículo).")


H3("Procedencia y construcción del dataset")
B("**Nombre y autores:** BRISC (BRain tumor Image Segmentation & Classification), de Fateh y colaboradores (2025), artículo arXiv 2506.14318. Se publica en Kaggle como `briscdataset/brisc2025`.")
B("**Qué dicen sus autores:** es una colección de cortes de resonancia **T1 con contraste**, en tres planos, con máscaras de segmentación a nivel de píxel **revisadas y corregidas por expertos médicos**, según su README. "
  "No pudimos comprobar esa revisión; la tomamos como declaración de los autores.")
B("**De dónde viene el material:** según el artículo, parte de las imágenes proviene de colecciones públicas previas (la de Nickparvar, que reúne los conjuntos de Cheng, SARTAJ y Br35H).")
B("**Dos afirmaciones de su README que al medir no se cumplen del todo:** dice que las clases están \"balanceadas\" (hay entre 1.207 y 1.757 imágenes por clase, una razón de 1,46 entre la mayor y la menor) y que las máscaras son \"binarias\" (todas traen valores intermedios).")
H3("Qué contiene cada archivo")
T([["Elemento", "Formato", "Cantidad", "Detalle medido"],
   ["Imágenes de clasificación", "JPEG", "6.000", "2.866 en escala de grises (1 canal) y 3.134 en RGB; en las 3.134 RGB los tres canales son iguales, por lo que **todas son en la práctica de un solo canal**"],
   ["Imágenes de segmentación", "JPEG", "4.793", "Repiten imágenes de la tarea de clasificación (el manifest muestra hashes repetidos entre las dos tareas)"],
   ["Máscaras", "PNG, 1 canal", "4.793", "Mismo nombre que su imagen y mismo tamaño; valores intermedios entre 0 y 6 y algunos 250"],
   ["Tamaño de archivo (imágenes de clasificación)", "—", "—", "Media por clase: glioma 19,1 KB · meningioma 22,4 KB · sin tumor 20,7 KB · hipófisis 29,0 KB (mínimo 4,9 KB; máximo 252,8 KB)"],
   ["Imágenes no cuadradas", "—", "659 de 6.000", "11 % de las imágenes; la mayoría de 512×512"],
   ["manifest.csv y manifest.json", "CSV / JSON", "15.586 filas", "16 campos por archivo (siguiente tabla) y sus sumas `.sha256`"]],
  [4.0, 2.2, 2.2, 8.2], size=8.5)
P("**Diccionario del manifest** (una fila por imagen y por máscara):", keep=True)
T([["Campo", "Qué contiene", "Ejemplo"],
   ["relative_path, filename", "Ruta relativa y nombre del archivo", "brisc2025_test_00001_gl_ax_t1.jpg"],
   ["task", "Tarea a la que pertenece el archivo", "classification / segmentation"],
   ["split", "Partición oficial", "train / test"],
   ["index", "Número correlativo de la imagen (no es ID de paciente)", "1"],
   ["tumor_code, tumor_label", "Clase en código y en texto", "gl / glioma; me; pi; no / no_tumor"],
   ["plane_code, plane_label", "Plano de corte", "ax / axial; co / coronal; sa / sagittal"],
   ["sequence", "Secuencia de resonancia (siempre T1)", "T1"],
   ["is_mask, linked_image", "Si el archivo es máscara y a qué imagen corresponde", "True / ruta de la imagen"],
   ["width, height", "Ancho y alto en píxeles", "512 / 512"],
   ["file_size_bytes, sha256", "Tamaño del archivo y huella SHA-256 (sirve para detectar duplicados exactos)", "25.826 / b513a83f…"]],
  [4.2, 8.0, 4.4], size=8.5)

H3("Contenido de las imágenes: intensidad y encuadre")
P("Para saber si las imágenes de las distintas clases son comparables se midió, en cada imagen, el brillo medio del cerebro (sin contar el fondo negro) y la fracción de la imagen que ocupa el cerebro "
  "(medido con el script `perfil_datos.py`). Los resultados por clase son distintos:", keep=True)
T([["Clase", "Imágenes", "En RGB (no en gris)", "Brillo medio del cerebro (0–255)", "Fracción de imagen con cerebro"],
   ["Glioma", "1.401", "12,4 %", "67,5 ± 9,9", "0,48"],
   ["Meningioma", "1.635", "56,6 %", "81,6 ± 15,2", "0,53"],
   ["Tumor de hipófisis", "1.757", "47,1 %", "70,5 ± 7,9", "0,68"],
   ["Sin tumor", "1.207", "100 %", "87,5 ± 17,8", "0,58"]],
  [3.6, 2.2, 3.2, 4.0, 3.6], size=8.5, align_right=(1, 2, 3, 4))
P("Por plano, la fracción de cerebro (mediana) es 0,46 en axial, 0,56 en coronal y 0,63 en sagital, y el brillo medio es 79,8 en axial y 74,4–74,5 en coronal y sagital.")
FIG(SAL + "perfil_datos/intensidad_y_cerebro_por_clase.png", 14.0, "Brillo medio del cerebro y fracción de la imagen ocupada por cerebro, por clase.")
FIG(FIGS + "modo_por_clase.png", 11.0, "Modo de archivo de las imágenes (escala de grises o RGB) por clase.")
P("**Cómo interpretarlo:**", keep=True)
B("**Todas las imágenes de \"sin tumor\" están en RGB y casi todas las de glioma en escala de grises.** Al cargar las imágenes se convierten a 3 canales y esa diferencia desaparece, "
  "pero indica que las clases pueden venir de colecciones o procesos distintos.")
B("El brillo y el encuadre **cambian con la clase** (y también con el plano). Es un riesgo de **atajo**: el modelo podría aprender de qué colección viene la imagen en lugar de reconocer el tumor, "
  "y rendir peor con imágenes de otro origen. Con estos datos no se puede separar el efecto de la clase del de la colección ni del plano; es una hipótesis para comprobar.")
B("**Qué se probará en el Sprint 2:** normalizar la intensidad por imagen, aumento de brillo y contraste, y reportar el desempeño por plano y por tipo de archivo (gris o RGB). "
  "Se propone registrar este riesgo como **R09** en el Registro de Riesgos.")

H3("Dónde está el tumor")
P("Se promedió, por clase, dónde cae cada máscara dentro de la imagen (rellenada a cuadrado, mezclando los tres planos):", keep=True)
FIG(SAL + "perfil_datos/mapa_localizacion_tumor.png", 15.0, "Figura 7. Proporción de máscaras que cubren cada zona de la imagen, por clase.")
T([["Clase", "Centro del tumor, horizontal (0 = izquierda, 1 = derecha)", "Centro del tumor, vertical (0 = arriba, 1 = abajo)"],
   ["Glioma", "0,485 ± 0,115", "0,427 ± 0,106"],
   ["Meningioma", "0,502 ± 0,126", "0,439 ± 0,149"],
   ["Tumor de hipófisis", "0,501 ± 0,052", "0,509 ± 0,079"]],
  [4.0, 6.3, 6.3], size=8.5, align_right=(1, 2))
B("Los tumores de hipófisis están **muy concentrados en el centro** de la imagen (desvío 0,05 horizontal): la posición es una pista fuerte para esa clase. Los gliomas y meningiomas se distribuyen por una zona más amplia.")
B("**Ninguna máscara toca el borde de la imagen:** no hay tumores cortados por el encuadre.")
B("Como las máscaras de los tres planos están mezcladas, el mapa muestra la tendencia general y no la anatomía de un plano concreto.")

H3("Regiones, tamaño y contraste del tumor")
T([["Aspecto", "Resultado"],
   ["Regiones separadas por máscara", "Una sola región en 4.394 máscaras (91,7 %); dos en 327; tres en 57; cuatro en 11; cinco en 4. **399 máscaras (8,3 %) tienen más de una región** (glioma 9,3 %, meningioma 9,1 %, hipófisis 6,8 %)"],
   ["Tamaño del tumor (percentiles 5 / 50 / 95 de la imagen)", "Glioma 0,29 / 1,52 / 4,81 % · meningioma 0,39 / 1,74 / 5,46 % · hipófisis 0,26 / 0,86 / 2,28 %"],
   ["Tumores extremos", "86 máscaras ocupan menos del 0,2 % de la imagen y 180 más del 5 %"],
   ["Tamaño por plano (mediana)", "Axial 1,20 % · coronal 1,37 % · sagital 1,15 %"],
   ["Contraste (brillo del tumor ÷ brillo del resto del cerebro)", "El tumor es más brillante que el resto en el 87,3 % de las máscaras. Mediana: glioma 1,07 (casi igual al tejido y con una parte más oscura), meningioma 1,52, hipófisis 1,45"]],
  [5.0, 11.6], size=8.5)
FIG(SAL + "perfil_datos/contraste_y_regiones.png", 14.0, "Figura 8. Contraste del tumor con el resto del cerebro y número de regiones separadas por máscara.")
P("**Implicaciones:** el tumor es pequeño y a veces está en varios pedazos, por lo que el modelo de segmentación debe poder marcar varias regiones; "
  "y los gliomas, con poco contraste respecto del tejido, serán probablemente los más difíciles de segmentar (hipótesis a comprobar).")

H3("Entrenamiento frente a prueba oficial: ¿se parecen?")
P("Se comparó la distribución de varias medidas entre el entrenamiento oficial y la prueba oficial, dentro de cada clase, con la prueba de Kolmogórov-Smirnov. "
  "El estadístico KS vale 0 si las dos distribuciones son iguales y 1 si no se parecen en nada; con miles de imágenes, hasta diferencias pequeñas dan un valor p muy bajo.", keep=True)
T([["Medida (clase)", "Entrenamiento", "Prueba", "KS", "Valor p"],
   ["Brillo medio del cerebro (meningioma)", "80,5", "86,5", "0,168", "< 0,001"],
   ["Fracción de cerebro (meningioma)", "0,521", "0,588", "0,192", "< 0,001"],
   ["Tamaño del tumor (meningioma)", "2,10 %", "2,87 %", "0,159", "< 0,001"],
   ["Tamaño del tumor (hipófisis)", "1,07 %", "0,89 %", "0,164", "< 0,001"],
   ["Tamaño de archivo (sin tumor)", "21,4 KB", "15,7 KB", "0,270", "< 0,001"],
   ["Tamaño de archivo (hipófisis)", "29,6 KB", "26,4 KB", "0,284", "< 0,001"],
   ["Brillo medio del cerebro (glioma)", "67,4", "68,0", "0,056", "0,50 (sin diferencia)"],
   ["Contraste del tumor (hipófisis)", "1,460", "1,451", "0,039", "0,83 (sin diferencia)"]],
  [6.4, 2.6, 2.4, 1.8, 3.4], size=8.5, align_right=(1, 2, 3))
FIG(FIGS + "train_vs_test.png", 15.0, "Entrenamiento frente a prueba oficial: imágenes en escala de grises y diferencia entre las dos particiones en cada medida.")
B("La prueba oficial **no es una copia exacta del entrenamiento**: tiene más imágenes en escala de grises (58,2 % frente a 45,7 %) y, en algunas clases, otro brillo, encuadre y tamaño de tumor. Las diferencias son moderadas.")
B("La validación que se separa del entrenamiento (sección 4.3) viene de la misma distribución que el entrenamiento, por lo que **no puede mostrar este desfase**. Es posible que el desempeño en la prueba oficial sea algo menor que en validación; "
  "conviene reportar los resultados también por clase, por plano y por tipo de archivo.")

H3("Resumen de la calidad de los datos")
T([["Dimensión", "Resultado"],
   ["Completitud", "6.000 imágenes y 4.793 máscaras; sin archivos faltantes; todas las imágenes con tumor tienen máscara; las imágenes sin tumor no tienen (esperado)"],
   ["Exactitud y consistencia", "0 archivos corruptos; cada máscara tiene el mismo tamaño que su imagen; nombres y códigos coherentes con las carpetas; 0 máscaras vacías"],
   ["Unicidad", "96 archivos con hash SHA-256 repetido (7 entre entrenamiento y prueba) y 113 a 123 imágenes de prueba casi idénticas a las de entrenamiento"],
   ["Formato", "Máscaras con valores intermedios; imágenes en dos modos (gris y RGB) y 183 tamaños distintos"],
   ["Balance", "Clases entre 1.207 y 1.757 imágenes (razón 1,46); tumor del 0,9 % al 1,7 % de la imagen en mediana"],
   ["Representatividad", "Solo T1 con contraste; sin ID de paciente; diferencias de brillo y encuadre entre clases y entre entrenamiento y prueba; la prevalencia de tumor (80 %) no es la de una clínica"]],
  [3.6, 13.0], size=8.5)

H2("4.3 HU03 · Preparar los datos y definir la partición")
H3("Integridad de los datos")
P("Se abrieron y verificaron las 6.000 imágenes y las 4.793 máscaras, y se comprobó que cada máscara tuviera el mismo tamaño que su imagen. "
  "Resultado: **0 archivos corruptos, 0 máscaras sin imagen, 0 imágenes de segmentación sin máscara, 0 máscaras vacías y 0 diferencias de tamaño**. "
  "La prueba oficial (1.000 imágenes) queda reservada: no se usa para entrenar ni para ajustar, y el notebook comprueba que no cambie.")
H3("Por qué una partición propia de validación")
P("BRISC solo trae entrenamiento (5.000) y prueba (1.000). En el Sprint 2 se compararán varias arquitecturas e hiperparámetros, y esas decisiones necesitan medirse con datos "
  "que el modelo no vio. Si se usara la prueba oficial para elegir, dejaría de ser una medida honesta del resultado final. Por eso la validación se separa del entrenamiento oficial, "
  "y la prueba se usa una sola vez, al final.")
H3("Agrupación de imágenes casi idénticas (hash perceptual)")
P("Como BRISC no trae ID de paciente, una partición aleatoria podría mandar cortes casi idénticos del mismo paciente a entrenamiento y a validación, inflando las métricas. "
  "Para reducir ese riesgo se calcula un **dHash de 256 bits** de cada imagen (diferencias entre píxeles vecinos de una versión reducida a 17×16), se agrupan las imágenes cuya "
  "distancia de Hamming es menor o igual que un umbral (algoritmo union-find) y cada grupo completo se asigna a un único lado de la partición.")
P("**El umbral no se fijó a ciegas.** La primera versión usaba un hash de 64 bits con umbral 6 y, al ejecutarla con los datos reales, formó un grupo de 1.933 imágenes distintas "
  "(inútil para partir los datos). Se midió entonces la distancia de cada imagen a su vecino más cercano dentro del entrenamiento (percentiles 5/25/50/75/95 = 21/31/40/55/77 bits de 256) "
  "y se compararon varios umbrales:", keep=True)
T([["Umbral (de 256 bits)", "Grupos", "Tamaño del mayor", "Grupos con más de 1 imagen"],
   ["8", "4.952", "3", "42"],
   ["**16 (elegido)**", "**4.938**", "**3**", "**55**"],
   ["24", "4.727", "6", "217"],
   ["32", "4.077", "11", "548"]],
  [4.6, 3.0, 4.0, 5.0], align_right=(1, 2, 3))
P("Con el umbral 16 las 5.000 imágenes de entrenamiento quedan en 4.938 grupos; 55 grupos tienen más de una imagen y el mayor tiene 3. "
  "Es un criterio conservador: agrupa solo imágenes casi idénticas.")
H3("Fuga entre las particiones oficiales")
P("También se midió, solo como información, cuántas imágenes de la **prueba oficial** tienen un vecino casi idéntico en el entrenamiento oficial: "
  "**113 de 1.000 según el dHash** (distancia ≤ 16) y **123 de 1.000 según la correlación de miniaturas de 32×32** (≥ 0,98). "
  "Esto no se puede corregir (la prueba oficial no se toca), pero hay que declararlo al reportar el resultado final: parte de la prueba se parece al entrenamiento.")
H3("Partición entrenamiento / validación")
P("Se usó `StratifiedGroupKFold` de scikit-learn (5 pliegues, uno de ellos como validación, semilla 42), estratificando por **clase y plano de corte** y respetando los grupos. "
  "El notebook comprueba con aserciones que ningún grupo quede repartido entre entrenamiento y validación y que la prueba oficial no haya cambiado.", keep=True)
T([["Clase", "Entrenamiento", "Validación", "Prueba oficial", "Total"],
   ["Glioma", "917", "230", "254", "1.401"],
   ["Meningioma", "1.063", "266", "306", "1.635"],
   ["Tumor de hipófisis", "1.166", "291", "300", "1.757"],
   ["Sin tumor", "853", "214", "140", "1.207"],
   ["**Total**", "**3.999**", "**1.001**", "**1.000**", "**6.000**"]],
  [4.6, 3.0, 3.0, 3.2, 2.8], align_right=(1, 2, 3, 4))
T([["Plano", "Entrenamiento", "Validación", "Prueba oficial"],
   ["Axial", "1.275", "320", "398"],
   ["Coronal", "1.341", "335", "305"],
   ["Sagital", "1.383", "346", "297"]],
  [4.6, 4.0, 4.0, 4.0], align_right=(1, 2, 3))
P("La proporción de validación es de 20 % en cada clase (0,200 a 0,201) y en cada plano (0,200 a 0,201).")
FIG(FIGS + "particion.png", 15.0, "Partición propia: imágenes por clase y reparto de los planos de corte en cada partición.")
H3("Preprocesamiento y cargador de datos")
P("El cargador (`BriscDataset`) devuelve la imagen normalizada (3 × 256 × 256), la clase, la máscara y un indicador `tiene_mascara`. Decisiones fijadas:", keep=True)
T([["Aspecto", "Decisión", "Motivo"],
   ["Tamaño de entrada", "256 × 256 (puede bajarse a 224 para transfer learning)", "Equilibrio entre detalle y memoria de GPU (6 GB)"],
   ["Imágenes no cuadradas", "Se rellenan con ceros a cuadrado, centradas, y luego se redimensionan", "Hay 183 tamaños distintos; así no se deforman"],
   ["Interpolación", "Bilineal para imágenes, *nearest* para máscaras", "Las máscaras no deben mezclar valores"],
   ["Máscaras", "Se binarizan con umbral > 127", "Traen valores intermedios por remuestreo"],
   ["Imágenes sin tumor", "Máscara vacía e indicador `tiene_mascara = False`", "No tienen máscara; así se puede ignorar su pérdida de segmentación"],
   ["Normalización", "Media (0,485; 0,456; 0,406) y desvío (0,229; 0,224; 0,225) de ImageNet", "Compatible con transfer learning"],
   ["Aumento de datos", "No se define aún", "Se ajusta en el Sprint 2 (HU08)"],
   ["Reproducibilidad", "Semilla 42; listas de partición guardadas en CSV", "Mismos conjuntos en todas las ejecuciones"]],
  [3.4, 7.2, 6.0], size=8.5)
P("El cargador se probó con un lote real de 8 imágenes (formas 8×3×256×256 para imágenes y 8×1×256×256 para máscaras) usando `DataLoader` de PyTorch.")
FIG(SAL + "lote_de_prueba.png", 13.0, "Figura 5. Lote de prueba producido por el cargador (imagen normalizada y máscara superpuesta; las imágenes sin tumor no muestran máscara).")
H3("Archivos generados")
B("`particion/particion_train.csv`, `particion_val.csv`, `particion_test.csv` y `particion_completa.csv` (archivo, ruta, clase, plano, partición, grupo y máscara).")
B("`particion/config_preparacion.json` (semilla, umbrales, tamaño de entrada, normalización, conteos y número de grupos).")
B("`eda_resumen.json`, `figuras_eda/` y `lote_de_prueba.png`, además de copias de los notebooks con todas sus salidas en `notebooks_ejecutados/`.")

H2("4.4 HU04 · Presentación del Sprint 1")
B("**Revisión del Sprint** (13 diapositivas, para el lunes 12 de octubre): meta del sprint, el proyecto, el dataset verificado, historias del sprint, demostración con las figuras reales "
  "(conteos, imagen y máscara, tamaño del tumor), hallazgos, riesgos y plan del Sprint 2.")
B("**Retrospectiva del Sprint** (7 diapositivas): plantilla del docente con las preguntas (qué aprendimos, qué estamos haciendo bien, qué podemos mejorar, personas y relaciones). "
  "Las respuestas corresponden al equipo y se registran en la reunión de retrospectiva.")
B("El **Informe de Estado** y la **Minuta de Reunión** del periodo también quedaron llenos con el avance del sprint; los pocos campos de datos que solo conoce el equipo (por ejemplo, fecha y lugar de la reunión o gastos registrados) los completa el equipo.")

# ================================================================ 5. HALLAZGOS
H1("5. Hallazgos y decisiones técnicas")
T([["Hallazgo (con datos reales)", "Decisión"],
   ["Los conteos por clase y partición coinciden con el artículo; los de plano difieren levemente", "Usar los conteos observados y documentar la diferencia"],
   ["Las 1.207 imágenes sin tumor no tienen máscara (4.793 pares)", "Entrenar y evaluar la segmentación solo con imágenes con tumor; en un modelo multitarea, máscara vacía y pérdida de segmentación ignorada"],
   ["Máscaras con valores intermedios (0–6 y 250/255)", "Binarizar con umbral > 127 en el análisis y en el cargador"],
   ["183 tamaños de imagen; no hay una única relación de aspecto", "Rellenar a cuadrado sin deformar y redimensionar a 256×256"],
   ["El tumor ocupa en mediana entre 0,9 % y 1,7 % de la imagen", "Considerar pérdidas robustas al desbalance en la segmentación (Dice o focal) y métricas por clase en el Sprint 2"],
   ["Sin ID de paciente; un hash de 64 bits con umbral 6 agrupaba mal (grupo de 1.933)", "dHash de 256 bits con umbral 16, elegido midiendo la distancia al vecino más cercano"],
   ["113–123 imágenes de prueba tienen un vecino casi idéntico en entrenamiento", "No tocar la prueba; reportar el resultado final con esta advertencia"],
   ["96 archivos con SHA-256 repetido (7 entre entrenamiento y prueba)", "Documentar; la agrupación por hash los mantiene del mismo lado dentro del entrenamiento"],
   ["Licencia CC BY 4.0", "Citar a los autores de BRISC y de las fuentes originales en el informe final"],
   ["Brillo y encuadre cambian con la clase; 100 % de \"sin tumor\" en RGB y 88 % de glioma en gris (posible atajo por colección de origen)", "Normalizar intensidad, aumentar brillo/contraste y reportar por plano y tipo de archivo; proponer el riesgo R09"],
   ["La prueba oficial difiere moderadamente del entrenamiento (brillo, encuadre, tamaño del tumor, más imágenes en gris)", "Reportar resultados por subgrupo y avisar que la validación propia no muestra este desfase"],
   ["8,3 % de las máscaras tiene más de una región; los gliomas tienen poco contraste con el tejido", "El modelo de segmentación debe poder marcar varias regiones; se espera que el glioma sea la clase más difícil"],
   ["El 80 % de las imágenes de BRISC tiene tumor; en una clínica serían pocas (prevalencia baja)", "Las metas son para la demostración; para un uso real habría que medir precisión y especificidad con la prevalencia real"],
   ["La clase sin tumor incluye cerebros sanos y lesiones no neoplásicas", "Mencionarlo al interpretar los errores de clasificación"]],
  [8.1, 8.5], size=8.5)

# ================================================================ 6. CALIDAD
H1("6. Calidad y verificación")
B("**Criterios de aceptación:** las 17 historias tienen criterios en Trello (la \"Q\" de CRISP-ML(Q)). Los de HU01, HU02 y HU03 están marcados como cumplidos, con comentarios que "
  "documentan la evidencia.")
B("**Desarrollo de los notebooks en dos pasos:** primero se probaron con un dataset sintético que imita la estructura descrita en el artículo y después se ejecutaron con BRISC real. "
  "La segunda ejecución destapó tres defectos que la prueba sintética no podía mostrar; se corrigieron y se volvió a ejecutar todo:")
B("Las máscaras no binarias se marcaban como \"problema\" en las 4.793 (ahora se informan y se binarizan).", 1)
B("El chequeo de duplicados del manifest contaba la tarea de segmentación (ahora filtra solo clasificación).", 1)
B("El umbral del hash perceptual era demasiado laxo (ahora se eligió midiendo las distancias).", 1)
B("**Comparación con el artículo:** conteos por clase, partición y pares de segmentación contrastados con las Tablas 1 y 2 del artículo de BRISC; las diferencias de plano quedaron registradas.")
B("**Revisión del Acta de Constitución:** se contrastó con el artículo y se corrigieron la licencia, las referencias de resultados y el presupuesto; la licencia se verificó en los metadatos de Kaggle.")
B("**Reproducibilidad:** los notebooks usan semilla fija y se ejecutan de punta a punta con `python ejecutar_sprint1.py`.")

# ================================================================ 7. CAMBIOS
H1("7. Cambios respecto al plan inicial")
T([["Cambio", "Detalle"],
   ["Presentación del Sprint 1", "Se movió del viernes 9 al lunes 12 de octubre por decisión del equipo; los Sprints 2 y 3 no cambian (23 y 30 de octubre)."],
   ["Dónde se ejecuta", "Los notebooks se pensaron para Kaggle Notebooks y el equipo decidió ejecutarlos **en local**. Se ajustaron scripts, READMEs, el Acta de Constitución, el Registro de Riesgos y las presentaciones. "
                        "El dataset se guarda fuera de OneDrive para no sincronizar unas 15.000 imágenes."],
   ["Estimaciones de cómputo", "El Acta y el Registro de Riesgos asumían GPU gratuita en la nube; ahora indican una GPU local (NVIDIA RTX 4050 de 6 GB) con Google Colab como respaldo."],
   ["Redacción de supuestos", "Tras ejecutar con datos reales, los supuestos \"a confirmar\" (imágenes sin máscara, licencia, partición) pasaron a hechos verificados en los documentos."]],
  [4.2, 12.4], size=9)

# ================================================================ 8. RIESGOS
H1("8. Riesgos y limitaciones")
P("Estado del Registro de Riesgos al cierre del Sprint 1 (la calificación de probabilidad e impacto es una propuesta del equipo):", keep=True)
T([["ID", "Riesgo", "Estado y novedad del Sprint 1"],
   ["R01", "Fuga entre pacientes: sin ID de paciente, una validación aleatoria puede dar métricas infladas", "Activo. Mitigado en la validación con grupos por hash; se detectó además que 113–123 imágenes de prueba se parecen a entrenamiento."],
   ["R02", "Cambio de dominio: peor desempeño con T2, sin contraste u otros centros", "Activo. BRISC solo contiene T1 con contraste; avisarlo en la aplicación."],
   ["R03", "Las imágenes sin tumor no tienen máscara", "Confirmado (0 de 1.207). Mitigación aplicada en el diseño de datos."],
   ["R04", "Arquitecturas inestables en el artículo (Xception, DenseNet y MobileNetV2)", "Activo. Partir de EfficientNet o ResNet y repetir entrenamientos con distintas semillas."],
   ["R05", "Modelos demasiado pesados para correr en el navegador", "Activo. Se decide en el Sprint 3 (modelos livianos o cuantizados; API gratuita si no alcanza)."],
   ["R06", "Falta de GPU o sobreajuste", "Activo. GPU local de 6 GB; PyTorch instalado es solo CPU y requiere la versión CUDA."],
   ["R07", "Retraso del cronograma (Sprint 2 es el más largo; margen final de un día)", "Activo. El Sprint 1 cerró según lo planificado."],
   ["R08", "Uso indebido de la demo como herramienta de diagnóstico", "Activo. Aviso visible de demostración académica."],
   ["R09 (propuesto)", "Atajos por la procedencia de las imágenes: el brillo, el encuadre y el modo de archivo cambian con la clase", "Nuevo en este informe; aún no está en el Registro de Riesgos. Mitigar con normalización, aumento de brillo y contraste y evaluación por subgrupos."]],
  [1.8, 6.6, 8.2], size=8.5)

# ================================================================ 9. ENTORNO
H1("9. Entorno de trabajo y reproducibilidad")
T([["Elemento", "Detalle"],
   ["Equipo de ejecución", "Windows 11, Python 3.13, GPU NVIDIA RTX 4050 de 6 GB (PyTorch instalado con soporte solo CPU)"],
   ["Bibliotecas del Sprint 1", "numpy, pandas, matplotlib, pillow, scikit-learn, jupyter, nbconvert, ipykernel y kaggle (ver `requirements.txt`)"],
   ["Dataset", "`C:\\Users\\<usuario>\\datasets\\brisc2025` (fuera de OneDrive); variable `BRISC_ROOT` para usar otra ruta"],
   ["Repositorio", "https://github.com/Calamar49/CNN-BRISC (documentos de gestión, notebooks y scripts; sin dataset ni salidas pesadas)"],
   ["Tablero Scrum", "Trello: Product Backlog, Sprint 1, Sprint 2, Sprint 3, En progreso y Terminado"]],
  [4.0, 12.6], size=9)
P("Para reproducir el Sprint 1, desde la carpeta `03. Implementacion/Sprint 1 - Notebooks`:", keep=True)
P("`pip install -r requirements-probado.txt`  →  `python descargar_brisc.py`  →  `python ejecutar_sprint1.py`  →  `python perfil_datos.py` (opcional)", align=WD_ALIGN_PARAGRAPH.CENTER)
P("La carpeta `salida_sprint1/` se regenera en cada ejecución y no se sube a Git. El dataset se descarga sin iniciar sesión en Kaggle.")

# ================================================================ 10. NO HECHO
H1("10. Alcance: lo que no se hizo en este sprint")
B("**No se entrenó ningún modelo** ni se calcularon métricas de desempeño; las metas del Acta (accuracy, recall e IoU) aún no se han medido.")
B("No se definió el aumento de datos (se ajusta en HU08) ni se eligió arquitectura ni framework de entrenamiento definitivo (decisión de HU05).")
B("**Pendientes menores de gestión:** completar los campos del equipo en la Retrospectiva, en el Informe de Estado y en la Minuta, y llenar el Cronograma detallado (archivo `.mpp` de la plantilla, que sigue vacío).")

# ================================================================ 11. SPRINT 2
H1("11. Plan del Sprint 2 y recomendaciones")
T([["Historia", "Fechas (Gantt)", "Puntos", "Depende de"],
   ["HU05 · Seleccionar arquitectura y entrenar el modelo base", "12 – 13 oct", "3", "HU03"],
   ["HU06 · Entrenar el clasificador de tipo de tumor", "14 – 18 oct", "8", "HU05"],
   ["HU07 · Entrenar el modelo de segmentación U-Net", "14 – 18 oct", "8", "HU05"],
   ["HU08 · Ajustar hiperparámetros y mejorar los modelos", "19 – 21 oct", "5", "HU06, HU07"],
   ["HU09 · Registrar y comparar experimentos", "21 – 22 oct", "3", "HU08"],
   ["HU10 · Presentación del Sprint 2", "23 oct", "2", "HU09"],
   ["HU11 · (Opcional) Modelo multitarea", "solo si el cronograma lo permite", "8", "HU09"]],
  [8.0, 4.4, 1.6, 2.6], size=9, align_right=(2,))
P("Recomendaciones para el Sprint 2 que salen de lo aprendido en el Sprint 1:", bold=True, keep=True)
N("**Instalar PyTorch con soporte CUDA** para usar la GPU (la instalación actual es solo CPU) y decidir el framework y el formato de exportación del modelo.")
N("**Partir de arquitecturas con respaldo** (EfficientNetB0 y ResNet50 en el artículo) y repetir los entrenamientos con 2 o 3 semillas por la inestabilidad reportada.")
N("**Segmentación solo con imágenes con tumor** (3.933 de entrenamiento oficial) y una pérdida adecuada al desbalance de píxeles (el tumor es del 1–2 % de la imagen).")
N("**Elegir en validación, evaluar una vez en la prueba.** Una vez elegida la mejor configuración se puede reentrenar con los 5.000 del entrenamiento oficial, "
  "fijando el número de épocas obtenido en validación, y evaluar una sola vez en la prueba oficial (HU12). Es opcional.")
N("**Guardar por experimento** la curva de entrenamiento, la época óptima, la semilla y las métricas por clase (HU09).")
N("**Reportar la advertencia de la prueba:** 113–123 imágenes de prueba tienen un vecino casi idéntico en entrenamiento.")

# ================================================================ ANEXOS
H1("Anexo A. Inventario de entregables")
T([["Carpeta", "Archivos"],
   ["01. Inicio", "Acta de Constitución, Backlog del Producto, Registro de Interesados, Planificación de Lanzamiento"],
   ["02. Planificacion", "Backlog Detallado del Producto, Backlog del Sprint (hoja Sprint 1), Cronograma de alto nivel – Hitos, Cronograma detallado (.mpp, sin llenar), "
                         "Presupuesto del Proyecto, Registro de Riesgos, CSV del Gantt"],
   ["03. Implementacion", "Informe de Estado del Proyecto, Minuta de Reunión, Registro de Impedimentos, Informe del Sprint 1 (este documento) y la carpeta "
                          "Sprint 1 - Notebooks (dos notebooks, descargar_brisc.py, ejecutar_sprint1.py, perfil_datos.py y README)"],
   ["04. Revision y retrospectiva", "Revisión del Sprint (13 diapositivas) y Retrospectiva del Sprint (7 diapositivas)"],
   ["05. Lanzamiento y 06. Cierre", "Plantillas sin llenar; se usan al final del proyecto"],
   ["Raíz del repositorio", "README.md, requirements.txt y .gitignore"]],
  [4.4, 12.2], size=9)

H1("Anexo B. Referencias")
B("Dataset BRISC 2025 (Kaggle): https://www.kaggle.com/datasets/briscdataset/brisc2025/ · Licencia CC BY 4.0 (metadatos de Kaggle).")
B("Artículo del dataset: https://arxiv.org/abs/2506.14318")
B("Repositorio del proyecto: https://github.com/Calamar49/CNN-BRISC")
B("Tablero Scrum (Trello): https://trello.com/b/swXjkoyS/product-backlog-proyecto-redes-neuronales-brisc")
B("Plantillas de gestión: repositorio del docente https://github.com/hespetia/gestion-proyectos")
B("Sector de salud peruano: Videnza Consultores para ComexPerú, Situación actual del mercado de salud peruano (agosto de 2022): https://www.comexperu.org.pe/upload/articles/publicaciones/situacion-actual-del-mercado-de-salud-peruano-informe.pdf")
B("Resonancias del Minsa (primer semestre): https://panamericana.pe/locales/452274-ministerio-salud-realizo-31-000-atenciones-modernos-resonadores-tomografos")
B("Referencia internacional de incidencia: CBTRUS Statistical Report (EE. UU.), cifras tomadas de extractos de búsqueda: https://pmc.ncbi.nlm.nih.gov/articles/PMC9113389")
B("Estadísticas de cáncer en Perú: IARC, Global Cancer Observatory, Cancer Today, estimaciones GLOBOCAN 2024 (ficha de Perú): https://gco.iarc.who.int/media/globocan/factsheets/populations/604-peru-fact-sheet.pdf")
B("Radiólogos, tecnólogos médicos y patólogos en Perú (Observatorio de Recursos Humanos en Salud del Minsa, citado por Panamericana): https://panamericana.pe/locales/459222-cuatro-regiones-cierran-2025-patologos-lima-concentra-casi-mitad-oncologos")

# orden de los hijos segun el esquema OOXML (Word es estricto con el orden)
_ORD = {
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
              "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"],
}
for _tag, _order in _ORD.items():
    for _el in d.element.body.iter(qn("w:" + _tag)):
        _kids = list(_el)
        _kids.sort(key=lambda k: _order.index(k.tag.split("}")[1]) if k.tag.split("}")[1] in _order else len(_order))
        for _k in _kids:
            _el.remove(_k)
        for _k in _kids:
            _el.append(_k)
for _z in d.settings.element.iter(qn("w:zoom")):
    if _z.get(qn("w:percent")) is None:
        _z.set(qn("w:percent"), "100")
d.core_properties.title = "Informe del Sprint 1 – Sistema Inteligente: tumores cerebrales en MRI"
d.core_properties.author = "Equipo del proyecto"
d.save(OUT)
print("ok", OUT)
