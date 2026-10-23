# -*- coding: utf-8 -*-
"""Presentaciones del Sprint 1 con diseño claro y formal (fondo blanco, 16:9).

Genera:  04. Revision y retrospectiva/Revisión del Sprint.pptx  y  Retrospectiva del Sprint.pptx
"""
from pathlib import Path as _P
_AQUI = _P(__file__).resolve().parent
_RAIZ = _AQUI.parent
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

ROOT = str(_RAIZ).replace("\\", "/") + "/"
SAL1 = ROOT + "03. Implementacion/Sprint 1 - Notebooks/salida_sprint1/"
FIG = SAL1 + "figuras_eda/"
PER = SAL1 + "perfil_datos/"
GF = str(_AQUI / "figs").replace("\\", "/") + "/"

def rgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

NAVY, INK, MUTED, LINE, PANEL, WHITE = rgb("1F3864"), rgb("263238"), rgb("6B7280"), rgb("D5DBE5"), rgb("F4F6FA"), rgb("FFFFFF")
STAGE = {0: rgb("6B7280"), 1: rgb("2B7A8B"), 2: rgb("2F5597"), 3: rgb("5E8C6A"), 4: rgb("6B7280")}
STAGE_TXT = {1: "ETAPA 1 · COMPRENSIÓN DEL NEGOCIO", 2: "ETAPA 2 · COMPRENSIÓN DE LOS DATOS",
             3: "ETAPA 3 · PREPARACIÓN DE LOS DATOS", 4: "CIERRE DEL SPRINT 1", 0: "SPRINT 1"}
FONT = "Calibri"
SW, SH = 13.333, 7.5
LM, CW = 0.7, 11.93          # margen izquierdo y ancho util
FOOT = "Sistema Inteligente: tumores cerebrales en MRI  ·  Revisión del Sprint 1"


# ------------------------------------------------------------------ primitivas
def run(p, text, size, bold=False, color=INK, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = FONT
    r.font.color.rgb = color
    return r


def text(slide, x, y, w, h, content, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    lines = content if isinstance(content, list) else [content]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        run(p, ln, size, bold, color, italic)
    return tb


def rect(slide, x, y, w, h, fill, line=None, kind=MSO_SHAPE.RECTANGLE, lw=0.75):
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
    return sp


def label_shape(slide, x, y, w, h, fill, txt, size=10, color=WHITE, bold=True, kind=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.CENTER):
    sp = rect(slide, x, y, w, h, fill, None, kind)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = 0.5
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run(p, txt, size, bold, color)
    return sp


def bullets(slide, items, x, y, w, h, size=16, accent=NAVY, gap=7):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
        sub = t.startswith("  ")
        run(p, ("–  " if sub else "■  "), size - (6 if not sub else 2), False, accent)
        run(p, t.strip(), size - (2 if sub else 0), False, INK)
        if sub:
            p.level = 1
    return tb


def pic(slide, path, x, y, w=None, h=None, border=True):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    pc = slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)
    if border:
        pc.line.color.rgb = LINE
        pc.line.width = Pt(0.75)
    return pc


def _cell_border(cell, color="D5DBE5", width=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    # los bordes van antes del relleno en el esquema
    first = tcPr[0] if len(tcPr) else None
    for tag, vis in (("a:lnL", False), ("a:lnR", False), ("a:lnT", False), ("a:lnB", True)):
        ln = etree.SubElement(tcPr, qn(tag))
        if vis:
            ln.set("w", str(width))
            sf = etree.SubElement(ln, qn("a:solidFill"))
            c = etree.SubElement(sf, qn("a:srgbClr"))
            c.set("val", color)
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        if first is not None:
            tcPr.remove(ln)
            first.addprevious(ln)


def table(slide, data, x, y, w, colw, size=13, rowh=0.5, head_fill=NAVY, first_col_bold=False, accent=None):
    shp = slide.shapes.add_table(len(data), len(data[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(data)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    for sid in tblPr.findall(qn("a:tableStyleId")):
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"       # sin estilo ni cuadricula
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for j, cw in enumerate(colw):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(data):
        tbl.rows[i].height = Inches(rowh)
        for j, t in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.12)
            c.margin_top = c.margin_bottom = Inches(0.05)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            is_head = i == 0
            run(p, t, size, is_head or (first_col_bold and j == 0), WHITE if is_head else INK)
            c.fill.solid()
            c.fill.fore_color.rgb = head_fill if is_head else (PANEL if i % 2 == 0 else WHITE)
            _cell_border(c)
    return tbl


# ------------------------------------------------------------------ estructura de diapositiva
BLOQUES = [("Joseph Coavoy", 8), ("Manuel Cuchuyrrumi", 7), ("Brayan Huallpatuiro", 7), ("Frank Mamani", 7)]


def expositor_por_diapositiva():
    """reparto equitativo y continuo: 8 + 7 + 7 + 7 diapositivas"""
    mapa, n = {}, 0
    for nombre, cant in BLOQUES:
        for _ in range(cant):
            n += 1
            mapa[n] = nombre
    return mapa


class Deck:
    def __init__(self, foot, presenters=None):
        self.presenters = presenters
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self.n = 0
        self.foot = foot

    def slide(self):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        return s

    def footer(self, s):
        rect(s, LM, 7.0, CW, 0.008, LINE)
        text(s, LM, 7.05, 8.0, 0.3, self.foot, 9, False, MUTED)
        if self.presenters and self.n in self.presenters:
            text(s, SW - LM - 4.6, 7.05, 3.5, 0.3, "Expone: " + self.presenters[self.n], 10, True, NAVY, PP_ALIGN.RIGHT)
        text(s, SW - LM - 0.9, 7.05, 0.9, 0.3, str(self.n), 9, False, MUTED, PP_ALIGN.RIGHT)

    def chrome(self, s, stage, title):
        col = STAGE[stage]
        rect(s, 0, 0, SW, 0.12, col)
        label_shape(s, LM, 0.36, 3.9, 0.3, col, STAGE_TXT[stage], 10)
        text(s, LM, 0.74, CW, 0.7, title, 28, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        rect(s, LM, 1.5, CW, 0.012, LINE)
        self.footer(s)

    def content(self, stage, title, items=None, size=18):
        s = self.slide()
        self.chrome(s, stage, title)
        if items:
            bullets(s, items, LM, 1.8, CW, 5.0, size, STAGE[stage])
        return s

    def divider(self, stage, big, title, sub, hizo, entregables, resultado):
        s = self.slide()
        col = STAGE[stage]
        rect(s, 0, 0, SW, 0.12, col)
        label_shape(s, LM, 0.36, 3.9, 0.3, col, STAGE_TXT[stage], 10)
        # columna izquierda: numeral y datos
        text(s, LM, 1.05, 3.2, 1.7, big, 96, True, col)
        text(s, LM, 2.75, 3.4, 1.3, title, 24, True, NAVY)
        text(s, LM, 4.15, 3.4, 1.4, sub, 14, False, MUTED, spacing=1.15)
        rect(s, 4.45, 1.0, 0.012, 5.8, LINE)
        # tarjetas
        x, w = 4.8, 7.83
        y = 1.0
        for head, items, size, h in (("QUÉ SE HIZO", hizo, 14, 3.05), ("ENTREGABLES", entregables, 14, 1.15), ("RESULTADO", resultado, 14, 1.4)):
            rect(s, x, y, 0.06, h, col)
            text(s, x + 0.2, y - 0.02, w - 0.2, 0.3, head, 11, True, col)
            bullets(s, items, x + 0.2, y + 0.3, w - 0.25, h - 0.3, size, col, gap=4)
            y += h + 0.17
        self.footer(s)
        return s

    def tile(self, s, x, y, w, h, big, small, col):
        rect(s, x, y, w, h, WHITE, LINE, lw=1)
        rect(s, x, y, 0.07, h, col)
        text(s, x + 0.25, y + 0.15, w - 0.4, 0.8, big, 40, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.25, y + 1.0, w - 0.4, h - 1.05, small, 13, False, MUTED, spacing=1.1)


# ------------------------------------------------------------------ Revision del Sprint
def build_review():
    D = Deck(FOOT, expositor_por_diapositiva())
    prs = D.prs

    # ---- portada
    s = D.slide()
    rect(s, 0, 0, 0.45, SH, rgb("1F3864"))
    text(s, 1.1, 0.9, 11, 0.4, "REUNIÓN DE DEMOSTRACIÓN Y REVISIÓN DEL SPRINT 1", 14, True, STAGE[1])
    text(s, 1.1, 1.5, 10.8, 2.2, "Sistema Inteligente para clasificar y segmentar tumores cerebrales en resonancias magnéticas", 38, True, NAVY, spacing=1.05)
    text(s, 1.1, 3.75, 10.8, 0.5, "desplegado en una aplicación web de apoyo académico", 20, False, MUTED)
    rect(s, 1.1, 4.5, 1.6, 0.04, STAGE[2])
    text(s, 1.1, 4.7, 11, 0.4, "Proyecto académico de Redes Neuronales · Ingeniería de Sistemas – Inteligencia Artificial", 14, False, INK)
    text(s, 1.1, 5.1, 11, 0.8,
         ["Coavoy Cruz Joseph Gabriel · Cuchuyrrumi Mamani Manuel Rodrigo",
          "Huallpatuiro Rafaile Brayan · Mamani Acuña Frank Joseph"], 14, False, INK, spacing=1.15)
    text(s, 1.1, 6.0, 11, 0.4, "Docente revisor: Ing. Hugo Espetia Huamanga  ·  12 de octubre de 2026", 14, False, MUTED)
    for i, (txt, st) in enumerate((("Etapa 1 · Negocio", 1), ("Etapa 2 · Datos", 2), ("Etapa 3 · Preparación", 3))):
        label_shape(s, 1.1 + i * 2.35, 6.6, 2.2, 0.3, STAGE[st], txt, 10)
    text(s, SW - 0.7 - 4.2, 6.62, 4.2, 0.3, "Abre la exposición: " + expositor_por_diapositiva()[1], 11, True, NAVY, PP_ALIGN.RIGHT)

    # ---- agenda
    s = D.slide()
    rect(s, 0, 0, SW, 0.12, STAGE[0])
    text(s, LM, 0.6, 5, 0.8, "Agenda", 34, True, NAVY)
    rect(s, LM, 1.5, CW, 0.012, LINE)
    rows = [("·", 0, "Meta del sprint y el proyecto", "Antes de las etapas"),
            ("1", 1, "Comprensión del negocio", "HU01 · 1 de octubre"),
            ("2", 2, "Comprensión de los datos", "HU02 · 2 a 5 de octubre"),
            ("3", 3, "Preparación de los datos", "HU03 · 6 a 8 de octubre"),
            ("·", 4, "Cierre: hallazgos, riesgos y plan del Sprint 2", "HU04 · 12 de octubre")]
    y = 1.85
    for num, st, ttl, sub in rows:
        sp = rect(s, LM, y, 0.62, 0.62, STAGE[st], None, MSO_SHAPE.OVAL)
        tf = sp.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run(p, num, 20, True, WHITE)
        text(s, LM + 0.9, y - 0.02, 9.5, 0.4, ttl, 22, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, LM + 0.9, y + 0.36, 9.5, 0.3, sub, 13, False, MUTED)
        y += 0.98
    D.footer(s)

    # ---- meta
    s = D.content(0, "Meta del Sprint 1")
    bullets(s, ["Entender el problema y el dataset BRISC a fondo",
                "Dejar los datos listos y particionados, sin fuga evidente",
                "Metodología híbrida: Scrum + CRISP-ML(Q)",
                "Trabajo del sprint: 1 al 8 de octubre · Presentación: 12 de octubre"], LM, 1.9, 6.6, 4.5, 20, STAGE[0], gap=14)
    rect(s, 7.8, 1.9, 4.83, 4.6, PANEL)
    text(s, 8.1, 2.05, 4.3, 0.35, "FASES DE CRISP-ML(Q) DEL SPRINT", 11, True, MUTED)
    for i, (t, st) in enumerate((("Comprensión del negocio", 1), ("Comprensión de los datos", 2), ("Preparación de los datos", 3))):
        rect(s, 8.1, 2.65 + i * 1.2, 0.07, 0.8, STAGE[st])
        text(s, 8.35, 2.65 + i * 1.2, 4.1, 0.8, t, 18, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.1, 6.0, 4.3, 0.4, "El modelado empieza en el Sprint 2", 12, False, MUTED, italic=True)

    # ---- proyecto
    s = D.content(0, "El proyecto: ¿para qué? ¿para dónde?")
    for i, (h, t, st) in enumerate((("¿PARA QUÉ?", "Mostrar, con fines académicos, cómo una red neuronal puede apoyar la interpretación de resonancias magnéticas cerebrales: identifica el tipo de tumor y delimita su contorno.", 1),
                                    ("¿PARA DÓNDE?", "Una aplicación web de demostración donde se sube una MRI y se obtiene la clase, la probabilidad y la máscara del tumor.", 2))):
        x = LM + i * 6.1
        rect(s, x, 1.9, 5.83, 3.6, PANEL)
        rect(s, x, 1.9, 5.83, 0.07, STAGE[st])
        text(s, x + 0.35, 2.2, 5.1, 0.4, h, 14, True, STAGE[st])
        text(s, x + 0.35, 2.7, 5.1, 2.7, t, 20, False, INK, spacing=1.15)
    rect(s, LM, 5.85, CW, 0.7, WHITE, LINE, lw=1)
    text(s, LM + 0.3, 5.85, CW - 0.6, 0.7, "No es una herramienta de diagnóstico: uso académico y de investigación.", 16, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)

    # ---- las 3 etapas
    s = D.content(0, "Las 3 etapas del Sprint 1")
    table(s, [["Etapa", "Historia", "Fechas", "Estado", "Entregable principal"],
              ["1 · Comprensión del negocio", "HU01", "1 oct", "Completada (acta firmada)", "Acta, backlog en Trello y Gantt"],
              ["2 · Comprensión de los datos", "HU02", "2 – 5 oct", "Completada", "Notebook, perfil de datos y figuras"],
              ["3 · Preparación de los datos", "HU03", "6 – 8 oct", "Completada", "Partición, cargador y scripts"],
              ["Cierre del sprint", "HU04", "12 oct", "Completada", "Esta presentación e informe"]],
          LM, 1.9, CW, [3.7, 1.2, 1.5, 2.6, 2.93], 15, 0.72, first_col_bold=True)
    bullets(s, ["Las etapas siguen las fases de CRISP-ML(Q); cada una tiene su historia de usuario en Trello",
                "En el Sprint 1 no se entrenó ningún modelo: eso es del Sprint 2"], LM, 5.8, CW, 1.1, 16, STAGE[0])

    # ============================================================ ETAPA 1
    D.divider(1, "1", "Comprensión del negocio", "Historia HU01\n1 de octubre\nFase 1 de CRISP-ML(Q)".replace("\n", " · ") if False else "Historia HU01 · 1 de octubre · Fase 1 de CRISP-ML(Q)",
              ["Definir el problema (¿para qué? ¿para dónde?) y el alcance",
               "Dimensionar el problema con datos públicos (GLOBOCAN, Minsa, SuSalud)",
               "Identificar a quién va dirigido y revisar el sector privado de salud con datos públicos",
               "Fijar objetivos y metas (accuracy ≥ 95 %, recall ≥ 90 %, IoU ≥ 70 %); evaluar factibilidad y riesgos",
               "Armar el Product Backlog (17 historias), el Gantt y las plantillas de gestión"],
              ["Acta de Constitución firmada · Product Backlog en Trello · Gantt y CSV · plantillas de gestión llenas"],
              ["Proyecto definido y factible con alcance acotado; 8 riesgos registrados; costo programado de S/ 0",
               "Proyecto académico: sin cliente, patrocinador ni Dueño del Producto"])

    s = D.content(1, "El problema en cifras")
    for i, (big, small) in enumerate((("1.616", "casos nuevos de cáncer de cerebro y sistema nervioso central en Perú (2024)"),
                                      ("1.147", "muertes por ese cáncer en Perú (2024)"),
                                      ("495", "médicos radiólogos en todo el país"),
                                      ("54,8 %", "de los radiólogos y tecnólogos médicos están en Lima"))):
        D.tile(s, LM + i * 3.04, 1.85, 2.85, 2.35, big, small, STAGE[1])
    bullets(s, ["La resonancia magnética es el estudio principal; el radiólogo decide si hay tumor, de qué tipo es y dónde está",
                "Pocos especialistas para toda la población y concentrados en Lima: esa es la motivación; el proyecto no mide ni demuestra una reducción de carga",
                "Fuentes: GLOBOCAN 2024 (IARC) y Observatorio de RR. HH. del Minsa, citado por Panamericana (nov. 2025)"],
            LM, 4.5, CW, 2.4, 15, STAGE[1])

    s = D.content(1, "El sector privado de salud")
    pic(s, GF + "sector_privado.png", 2.07, 1.7, w=9.2)
    bullets(s, ["52 % de los 19.067 establecimientos son privados, casi todos consultorios (9.508); ≈ 312 clínicas y hospitales de 2.º y 3.er nivel (cálculo propio)",
                "Seguro privado: 6,7 % del país y 14,5 % en Lima Metropolitana (2022)",
                "No hay cifras públicas de resonadores privados ni de pacientes por clínica: no se inventaron"],
            LM, 5.45, CW, 1.5, 15, STAGE[1], gap=4)

    s = D.content(1, "¿A quién está dirigido?")
    table(s, [["Grupo", "Relación con el proyecto", "Qué necesita"],
              ["Equipo del proyecto (4 integrantes)", "Usuario real de la demostración", "Aprender y demostrar un sistema de IA de punta a punta"],
              ["Docente revisor", "Evalúa el proyecto en cada sprint; no es patrocinador", "Evidencia del proceso y resultados medidos con rigor"]],
          LM, 1.85, CW, [3.9, 4.2, 3.83], 16, 0.85, first_col_bold=True)
    bullets(s, ["Proyecto académico: no tiene cliente, patrocinador ni Dueño del Producto; el equipo define y prioriza el backlog",
                "Uso académico y de investigación, con imágenes T1 con contraste",
                "No es una herramienta de diagnóstico y no se usa con pacientes"], LM, 4.8, CW, 2.1, 18, STAGE[1], gap=12)

    s = D.content(1, "Objetivos, metas y factibilidad")
    table(s, [["Nivel", "Objetivo", "Criterio de éxito"],
              ["Proyecto", "Sistema de IA de punta a punta con Scrum + CRISP-ML(Q)", "3 presentaciones (12, 23 y 30 oct) y entrega antes del 2 nov"],
              ["Demostración", "App web: clase, probabilidad y máscara", "Demo funcional con modelos exportados"],
              ["Clasificación", "4 clases", "Accuracy ≥ 95 % y recall ≥ 90 % por clase (prueba oficial)"],
              ["Segmentación", "Contorno del tumor", "IoU ≥ 70 % (prueba oficial, solo con tumor)"]],
          LM, 1.8, CW, [2.3, 4.7, 4.93], 14, 0.62, first_col_bold=True)
    bullets(s, ["Metas propuestas por el equipo con referencia al artículo (no son resultados propios); aún no se evaluaron",
                "Factible con alcance acotado: datos suficientes, GPU local de 6 GB, 33 días y modelo multitarea opcional",
                "Costo programado: S/ 0"], LM, 5.2, CW, 1.7, 16, STAGE[1])

    # ============================================================ ETAPA 2
    D.divider(2, "2", "Comprensión de los datos", "Historia HU02 · 2 a 5 de octubre · Fase 2 de CRISP-ML(Q)",
              ["Descargar BRISC 2025 y verificar su estructura, el manifest y la licencia",
               "Contar imágenes por clase, partición y plano de corte, y compararlas con el artículo",
               "Analizar las máscaras, los tamaños de imagen y los duplicados exactos (SHA-256)",
               "Perfilar el contenido: brillo, encuadre, ubicación, tamaño, regiones y contraste del tumor",
               "Comparar entrenamiento con prueba oficial y registrar la calidad y las limitaciones"],
              ["Notebook 01_HU02 y script perfil_datos.py · figuras y eda_resumen.json · informe con los resultados"],
              ["Conteos iguales a los del artículo; 4.793 pares imagen-máscara; sin tumor no tiene máscara",
               "Hallazgos: máscaras no binarias, 183 tamaños, atajos por procedencia, prueba distinta al entrenamiento, 96 duplicados"])

    s = D.content(2, "BRISC 2025 (verificado)")
    table(s, [["Clase", "Entrenamiento", "Prueba", "Total"],
              ["Glioma", "1.147", "254", "1.401"],
              ["Meningioma", "1.329", "306", "1.635"],
              ["Tumor de hipófisis", "1.457", "300", "1.757"],
              ["Sin tumor", "1.067", "140", "1.207"],
              ["Total", "5.000", "1.000", "6.000"]],
          LM, 1.85, 6.6, [2.5, 1.9, 1.1, 1.1], 15, 0.55, first_col_bold=True)
    bullets(s, ["Imágenes T1 con contraste en cortes axiales, coronales y sagitales",
                "Autores: Fateh et al. (2025)",
                "4.793 pares imagen-máscara (solo las 3 clases con tumor)",
                "Sin ID de paciente: riesgo de fuga entre particiones",
                "Kaggle (briscdataset/brisc2025) · licencia CC BY 4.0"], 7.6, 1.85, 5.0, 4.8, 16, STAGE[2], gap=10)

    s = D.content(2, "Qué contiene cada archivo")
    table(s, [["Elemento", "Formato", "Cantidad", "Detalle medido"],
              ["Imágenes de clasificación", "JPEG", "6.000", "2.866 en grises y 3.134 en RGB con canales iguales (en la práctica, un canal)"],
              ["Imágenes de segmentación", "JPEG", "4.793", "Repiten imágenes de clasificación"],
              ["Máscaras", "PNG", "4.793", "Valores intermedios (0 a 6 y algunos 250): se binarizan > 127"],
              ["Tamaños", "—", "183 combinaciones", "512×512 en 4.881 imágenes; 659 no son cuadradas"],
              ["manifest.csv / .json", "CSV / JSON", "15.586 filas", "16 campos por archivo, con SHA-256 para detectar duplicados"]],
          LM, 1.8, CW, [3.1, 1.4, 2.0, 5.43], 13, 0.66, first_col_bold=True)
    bullets(s, ["El README del dataset dice \"clases balanceadas\" y \"máscaras binarias\": al medir, las clases van de 1.207 a 1.757 imágenes y las máscaras no son estrictamente binarias"],
            LM, 6.0, CW, 0.9, 14, STAGE[2])

    s = D.content(2, "Conteos por clase y plano de corte")
    pic(s, FIG + "conteo_clase_particion.png", LM, 1.8, w=5.85)
    pic(s, FIG + "planos_por_clase.png", LM + 6.08, 1.8, w=5.85)
    bullets(s, ["Los conteos por clase y partición son idénticos a los del artículo",
                "Los conteos por plano difieren levemente de su Tabla 2 (por ejemplo, entrenamiento axial 1.595 frente a 1.550); se usan los observados"],
            LM, 5.2, CW, 1.7, 17, STAGE[2], gap=8)

    s = D.content(2, "Imagen y máscara por clase")
    pic(s, FIG + "galeria_imagen_mascara.png", LM, 1.7, h=5.15)
    bullets(s, ["Máscara del tumor en rojo sobre la imagen",
                "Fila inferior: clase sin tumor, que no tiene máscara",
                "Los tumores son pequeños y de formas muy variadas",
                "Se ven los tres planos de corte: axial, coronal y sagital"], 5.4, 1.9, 7.2, 4.5, 17, STAGE[2], gap=12)

    s = D.content(2, "Modo de archivo, brillo y encuadre")
    pic(s, GF + "modo_por_clase.png", LM, 1.75, w=5.6)
    pic(s, PER + "intensidad_y_cerebro_por_clase.png", 6.5, 1.75, w=6.1)
    bullets(s, ["100 % de \"sin tumor\" está en RGB y 88 % de glioma en grises",
                "El brillo y el encuadre cambian con la clase y con el plano",
                "Riesgo de atajo: el modelo podría aprender la colección de origen (propuesto como R09)"],
            LM, 4.9, CW, 2.0, 18, STAGE[2], gap=9)

    s = D.content(2, "Dónde está el tumor")
    pic(s, PER + "mapa_localizacion_tumor.png", 1.87, 1.7, w=9.6)
    bullets(s, ["Hipófisis muy concentrada en el centro de la imagen (desvío 0,05): la posición es una pista fuerte",
                "Glioma y meningioma ocupan una zona más amplia; ninguna máscara toca el borde de la imagen"],
            LM, 5.3, CW, 1.6, 17, STAGE[2], gap=8)

    s = D.content(2, "Tamaño del tumor")
    pic(s, FIG + "tamano_tumor.png", LM, 1.75, w=7.4)
    bullets(s, ["En mediana, el tumor ocupa entre 0,9 % (hipófisis) y 1,7 % (meningioma) de la imagen: el fondo domina",
                "Esto condiciona la función de pérdida de la segmentación en el Sprint 2"], 8.4, 1.85, 4.2, 4.8, 16, STAGE[2], gap=12)

    s = D.content(2, "Regiones y contraste")
    pic(s, PER + "contraste_y_regiones.png", 1.87, 1.7, w=9.6)
    bullets(s, ["8,3 % de las máscaras tiene más de una región: el modelo debe poder marcar varias",
                "El tumor es más brillante que el resto del cerebro en el 87 % de las máscaras",
                "Glioma: contraste mediano 1,07 (casi igual al tejido); se espera que sea la clase más difícil de segmentar"],
            LM, 5.4, CW, 1.55, 15, STAGE[2], gap=4)

    s = D.content(2, "Entrenamiento frente a prueba oficial")
    pic(s, GF + "train_vs_test.png", 2.27, 1.7, w=8.8)
    bullets(s, ["La prueba tiene más imágenes en grises (58,2 % frente a 45,7 %) y, en algunas clases, otro brillo, encuadre y tamaño de tumor",
                "La validación propia sale del entrenamiento y no muestra este desfase: reportar resultados por clase, plano y tipo de archivo"],
            LM, 5.7, CW, 1.25, 14, STAGE[2], gap=3)

    s = D.content(2, "Calidad de los datos")
    table(s, [["Dimensión", "Resultado"],
              ["Completitud", "6.000 imágenes y 4.793 máscaras; todas las imágenes con tumor tienen máscara"],
              ["Exactitud", "0 archivos corruptos; máscara e imagen del mismo tamaño; 0 máscaras vacías"],
              ["Unicidad", "96 archivos con SHA-256 repetido; 113 a 123 imágenes de prueba casi idénticas al entrenamiento"],
              ["Formato", "Máscaras con valores intermedios; 2 modos de imagen; 183 tamaños"],
              ["Balance", "Clases de 1.207 a 1.757 imágenes; tumor de 0,9 % a 1,7 % de la imagen"],
              ["Representatividad", "Solo T1 con contraste; sin ID de paciente; 80 % de las imágenes tiene tumor"]],
          LM, 1.8, CW, [2.8, 9.13], 14, 0.66, first_col_bold=True)

    # ============================================================ ETAPA 3
    D.divider(3, "3", "Preparación de los datos", "Historia HU03 · 6 a 8 de octubre · Fase 3 de CRISP-ML(Q)",
              ["Verificar la integridad: archivos corruptos y pares imagen-máscara",
               "Agrupar imágenes casi idénticas con un hash perceptual, eligiendo el umbral con datos reales",
               "Partir entrenamiento y validación sin tocar la prueba oficial (20 %, por clase y plano, semilla 42)",
               "Fijar el preprocesamiento y escribir el cargador de datos; probarlo con un lote real",
               "Guardar la partición y dejar todo reproducible con scripts"],
              ["Notebook 02_HU03 · listas de partición y config_preparacion.json · cargador de datos · descargar_brisc.py y ejecutar_sprint1.py"],
              ["3.999 entrenamiento, 1.001 validación y 1.000 prueba oficial intacta; 4.938 grupos de imágenes parecidas",
               "0 archivos corruptos; cargador probado con un lote; resultados idénticos al repetir la ejecución"])

    s = D.content(3, "Integridad y agrupación por hash")
    table(s, [["Umbral (de 256 bits)", "Grupos", "Mayor grupo", "Grupos con más de 1 imagen"],
              ["8", "4.952", "3", "42"],
              ["16 (elegido)", "4.938", "3", "55"],
              ["24", "4.727", "6", "217"],
              ["32", "4.077", "11", "548"]],
          LM, 1.8, 7.3, [2.4, 1.3, 1.6, 2.0], 15, 0.55, first_col_bold=True)
    bullets(s, ["Primer intento: hash de 64 bits con umbral 6 formó un grupo de 1.933 imágenes distintas; se midieron las distancias reales y se corrigió",
                "Integridad: 0 archivos corruptos, 0 máscaras sin imagen, 0 vacías y 0 con tamaño distinto",
                "Control informativo: 113 a 123 imágenes de prueba tienen un vecino casi idéntico en entrenamiento (no se corrige: la prueba no se toca)"],
            LM, 4.85, CW, 2.1, 16, STAGE[3])
    rect(s, 8.3, 1.8, 4.33, 2.7, PANEL)
    rect(s, 8.3, 1.8, 0.07, 2.7, STAGE[3])
    text(s, 8.6, 1.9, 3.9, 0.4, "POR QUÉ SE AGRUPAN", 11, True, STAGE[3])
    text(s, 8.6, 2.3, 3.9, 2.1, "BRISC no trae ID de paciente. Si cortes casi idénticos del mismo paciente quedaran en entrenamiento y en validación, las métricas saldrían infladas.",
         16, False, INK, spacing=1.15)

    s = D.content(3, "Partición propia de los datos")
    pic(s, GF + "particion.png", 2.07, 1.7, w=9.2)
    bullets(s, ["3.999 entrenamiento, 1.001 validación (20 %, en cada clase y plano) y 1.000 prueba oficial intacta",
                "Cada grupo de imágenes parecidas va completo a un solo lado",
                "La validación sirve para elegir modelos; la prueba se usa una sola vez, al final"],
            LM, 5.45, CW, 1.5, 15, STAGE[3], gap=4)

    s = D.content(3, "Preprocesamiento y cargador de datos")
    pic(s, SAL1 + "lote_de_prueba.png", LM, 1.75, w=7.5)
    bullets(s, ["Relleno a cuadrado sin deformar y 256×256",
                "Máscaras con umbral > 127",
                "Normalización de ImageNet",
                "Sin tumor: máscara vacía",
                "Probado con un lote real (8×3×256×256)"], 8.5, 1.85, 4.1, 4.8, 16, STAGE[3], gap=10)

    # ============================================================ CIERRE
    D.content(4, "Demostración del trabajo completado", [
        "Etapa 1: tablero de Trello con 17 historias, criterios de aceptación y fases CRISP-ML(Q)",
        "Etapa 2: notebook HU02 y perfil de datos con los datos reales de BRISC",
        "Etapa 3: notebook HU03 con integridad, hash, partición y cargador",
        "Los conteos por clase y partición coinciden con el artículo (diferencia 0 en las 8 combinaciones)",
    ], 22)

    s = D.content(4, "Hallazgos y decisiones")
    table(s, [["Etapa", "Hallazgo", "Decisión"],
              ["2", "Sin tumor no tiene máscara", "Segmentación solo con las 4.793 imágenes con tumor"],
              ["2 → 3", "Máscaras con valores intermedios; 183 tamaños", "Binarizar > 127; rellenar a cuadrado y redimensionar a 256"],
              ["2", "Tumor pequeño (0,9 % a 1,7 %) y a veces en varias regiones", "Pérdida robusta al desbalance (Dice o focal)"],
              ["2", "Brillo y encuadre cambian con la clase", "Normalizar, aumentar brillo y contraste; evaluar por subgrupo (R09)"],
              ["2", "Prueba distinta al entrenamiento; 113 a 123 imágenes parecidas", "No tocar la prueba; reportar por subgrupo y avisar"],
              ["3", "Hash de 64 bits agrupaba mal (grupo de 1.933)", "dHash de 256 bits con umbral 16, elegido midiendo distancias"]],
          LM, 1.75, CW, [1.1, 5.2, 5.63], 13, 0.62, first_col_bold=True)

    D.content(4, "Riesgos y limitaciones detectadas", [
        "Fuga entre pacientes: BRISC no trae ID de paciente; el resultado final se reporta solo con la prueba oficial y se informa que parte de ella se parece al entrenamiento",
        "Cambio de dominio: solo hay T1 con contraste",
        "Atajos por la procedencia de las imágenes (R09, propuesto): brillo, encuadre y modo de archivo cambian con la clase",
        "Imágenes sin tumor: no tienen máscara; la segmentación se entrena y evalúa solo con las 4.793 imágenes con tumor",
        "Arquitecturas inestables en el artículo: Xception, DenseNet y MobileNetV2",
        "Licencia: CC BY 4.0 (metadatos de Kaggle); hay que citar a los autores",
    ], 20)

    D.content(4, "Pendientes y plan del Sprint 2", [
        "Sprint 1 concluido: las 4 historias (HU01 a HU04) están terminadas y el acta está firmada",
        "Decisión pendiente: framework de entrenamiento y formato de exportación",
        "Instalar PyTorch con soporte CUDA para usar la GPU (la instalación actual es solo CPU)",
        "Sprint 2 (10–22 oct) – Modelado, presentación el 23 oct:",
        "  HU05 Arquitectura y modelo base (12–13 oct)",
        "  HU06 y HU07 Entrenar clasificador y U-Net (14–18 oct)",
        "  HU08 Ajuste de hiperparámetros (19–21 oct)",
        "  HU09 Registro y comparación de experimentos (21–22 oct)",
    ], 21)

    f = ROOT + "04. Revision y retrospectiva/Revisión del Sprint.pptx"
    prs.save(f)
    print("Revisión del Sprint:", len(prs.slides), "diapositivas")


# ------------------------------------------------------------------ Retrospectiva del Sprint
def build_retro():
    D = Deck("Sistema Inteligente: tumores cerebrales en MRI  ·  Retrospectiva del Sprint 1")
    prs = D.prs
    s = D.slide()
    rect(s, 0, 0, 0.45, SH, rgb("1F3864"))
    text(s, 1.1, 1.4, 11, 0.4, "REUNIÓN DEL EQUIPO", 14, True, STAGE[3])
    text(s, 1.1, 1.95, 10.8, 1.6, "Retrospectiva del Sprint 1", 44, True, NAVY)
    text(s, 1.1, 3.55, 10.8, 0.6, "Sistema Inteligente para clasificar y segmentar tumores cerebrales en MRI", 20, False, MUTED)
    rect(s, 1.1, 4.4, 1.6, 0.04, STAGE[2])
    text(s, 1.1, 4.65, 11, 0.4, "12 de octubre de 2026 · Scrum Master: Cuchuyrrumi Mamani Manuel Rodrigo", 14, False, INK)

    s = D.slide()
    rect(s, 0, 0, SW, 0.12, STAGE[0])
    text(s, LM, 0.6, 8, 0.8, "Temas de la retrospectiva", 32, True, NAVY)
    rect(s, LM, 1.5, CW, 0.012, LINE)
    temas = [("¿Qué aprendimos?", "Sobre el problema, los datos, las herramientas y la metodología"),
             ("¿Qué estamos haciendo bien?", "Lo que queremos mantener en el Sprint 2"),
             ("¿Qué podemos hacer mejor?", "Acciones concretas para el Sprint 2"),
             ("Personas", "Reparto del trabajo, apoyo entre compañeros y disponibilidad"),
             ("Relaciones", "Comunicación del equipo y con el docente")]
    y = 1.85
    for i, (t, sub) in enumerate(temas, 1):
        sp = rect(s, LM, y, 0.55, 0.55, STAGE[2], None, MSO_SHAPE.OVAL)
        tf = sp.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run(p, str(i), 18, True, WHITE)
        text(s, LM + 0.85, y - 0.04, 10.5, 0.4, t, 20, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, LM + 0.85, y + 0.33, 10.5, 0.3, sub, 13, False, MUTED)
        y += 0.98
    rect(s, LM, 7.0, CW, 0.008, LINE)
    text(s, LM, 7.05, 9.0, 0.3, D.foot, 9, False, MUTED)
    text(s, SW - LM - 1.0, 7.05, 1.0, 0.3, str(D.n), 9, False, MUTED, PP_ALIGN.RIGHT)

    for i, (t, sub) in enumerate(temas, 1):
        s = D.slide()
        col = STAGE[2]
        rect(s, 0, 0, SW, 0.12, col)
        label_shape(s, LM, 0.36, 2.6, 0.3, col, f"RETROSPECTIVA · TEMA {i}", 10)
        text(s, LM, 0.74, CW, 0.7, t, 28, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        rect(s, LM, 1.5, CW, 0.012, LINE)
        text(s, LM, 1.75, CW, 0.5, sub, 18, False, MUTED)
        rect(s, LM, 2.5, CW, 4.2, PANEL, LINE, lw=1)
        text(s, LM + 0.4, 2.8, CW - 0.8, 0.5, "[completar con el equipo en la retrospectiva]", 18, False, MUTED, italic=True)
        rect(s, LM, 7.0, CW, 0.008, LINE)
        text(s, LM, 7.05, 9.0, 0.3, D.foot, 9, False, MUTED)
        text(s, SW - LM - 1.0, 7.05, 1.0, 0.3, str(D.n), 9, False, MUTED, PP_ALIGN.RIGHT)

    f = ROOT + "04. Revision y retrospectiva/Retrospectiva del Sprint.pptx"
    prs.save(f)
    print("Retrospectiva:", len(prs.slides), "diapositivas")


if __name__ == "__main__":
    build_review()
    build_retro()
